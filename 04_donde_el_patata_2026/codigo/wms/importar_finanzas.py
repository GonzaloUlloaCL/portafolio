#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
importar_finanzas.py — Carga el HISTÓRICO de la planilla real al libro de caja.
================================================================================
Lee `planilla_2026_mejorada.xlsx` (hojas MAYO/JUNIO/JULIO) y carga en el módulo
de finanzas (wms/finanzas.py) los datos diarios: ventas por medio de pago
(snapshot), boleta manual, pago al delivery y egresos por categoría.

CÓMO LEE LAS CELDAS
  - Usa los VALORES cacheados (data_only=True): las celdas "concatenadas" de la
    planilla son fórmulas `=39300+14400+3000` cuyo resultado ya está calculado.
  - Celdas de texto tipo 'deli 25500' → extrae y suma los números; 'ejemplo' → 0.

MAPEO (template fijo de la planilla, idéntico en las 3 hojas)
  - Ingresos por FILA fija (las etiquetas de ingresos son inconsistentes entre
    hojas): r4=Pago al delivery, r5=Efectivo, r6=Transferencia, r7=POS, r9=Boleta
    Manual.  - Egresos: fila fija → categoría (r15..r59 + r65/r66).
  - El proveedor (col A, incl. comodines Otro/oferta/vacío) va en la NOTA del
    egreso; proveedor_id queda NULL (no se duplica el maestro del WMS).

LÍMITE HONESTO: la planilla NO marca factura sí/no por línea (solo totales del día
r76/r77), así que el histórico se importa con con_factura=0. El detalle de factura
por transacción empieza a existir desde la carga en vivo (POS/WhatsApp/SII).

Idempotente: re-importar un mes borra primero lo previo con origen='import' de ese
mes. NO toca datos cargados en vivo (origen manual/pos/sii/whatsapp).

Uso:
  python wms/importar_finanzas.py --dry-run                 # solo valida vs la planilla
  python wms/importar_finanzas.py --commit --db <ruta.db>   # importa a esa base
  python wms/importar_finanzas.py --commit                  # importa a pos/pedidos.db (REAL)
"""
import argparse
import calendar
import re
import sys
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
POS = RAIZ / "pos"
sys.path.insert(0, str(POS))
sys.path.insert(0, str(AQUI))
from nucleo import NucleoPedidos          # noqa: E402
from finanzas import NucleoFinanzas        # noqa: E402

PLANILLA_DEFAULT = Path(r"C:\Users\gulloa\Downloads\planilla_2026_mejorada.xlsx")
ANIO = 2026
MESES = {"MAYO": 5, "JUNIO": 6, "JULIO": 7}
COL_DIA1 = 3                                # columna C = día 1

# Filas de ingreso (fijas en las 3 hojas).
R_DELIVERY, R_EFECTIVO, R_TRANSFER, R_POS, R_BOLETA = 4, 5, 6, 7, 9

# Fila -> categoría de egreso (template de la planilla). Tomates/Paltas tienen 2
# filas (proveedor distinto) que caen en la misma categoría: es la M:N esperada.
FILA_CATEGORIA = {
    15: "Churrasco", 16: "Vienesas", 17: "Pan", 18: "Tomates", 19: "Tomates",
    20: "Paltas", 21: "Paltas", 22: "Empanadita Queso", 23: "Papas congeladas",
    24: "Huevos", 25: "Cartones, art de aseo", 26: "Aceite", 27: "Queso",
    28: "Sal, tapapecho, té", 29: "entrada", 30: "Gas", 33: "envases, bolsas",
    34: "Poroto verde", 35: "Toalla papel, servilletas", 36: "Vasos, huevo en polvo",
    37: "Ketchup, mostaza", 38: "Ají crema, salsa queso", 39: "bbq, alitas de pollo",
    40: "tiritas de pollo", 41: "nugget, aros de cebolla", 42: "Fruna 500cc",
    43: "Lata 350cc", 44: "Botella 3000cc", 46: "Patata 50 cuotas", 47: "Sueldo Chris",
    48: "Sueldo Yocy", 49: "Dividendo Huechuraba", 50: "Agua", 51: "Luz",
    52: "Internet", 53: "Teléfono celular", 54: "IVA", 55: "Patente", 56: "Bencina",
    57: "previred.com", 58: "contador", 59: "contribuciones Huechuraba",
    65: "Reparaciones", 66: "Compra de utensilios",
}
COL_MES = 34                               # columna AH = total del mes (para validar)


def num_celda(v):
    """Valor numérico de una celda: número directo, o suma de los números que
    contenga si es texto/fórmula ('=39300+14400' -> 53700; 'deli 25500' -> 25500;
    'ejemplo' -> 0)."""
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v)
    nums = re.findall(r"\d+", s.replace(".", "").replace(",", ""))
    return sum(int(x) for x in nums) if nums else 0


def _abrir_planilla(ruta):
    try:
        import openpyxl
    except ImportError:
        sys.exit("Falta openpyxl. Instala con: pip install openpyxl")
    if not Path(ruta).exists():
        sys.exit(f"No encontré la planilla: {ruta}")
    return openpyxl.load_workbook(ruta, data_only=True)


def _split_num_txt(vals):
    """Separa una lista de valores de celda en (suma de celdas numéricas, suma de
    números embebidos en celdas de TEXTO). La fórmula SUM de Excel solo cuenta las
    numéricas; el texto ('deli 25500') lo captura el importador pero la planilla no."""
    num = txt = 0
    for v in vals:
        if isinstance(v, (int, float)):
            num += int(round(v))
        elif isinstance(v, str):
            nums = re.findall(r"\d+", v.replace(".", "").replace(",", ""))
            txt += sum(int(x) for x in nums)
    return num, txt


def validar(ruta=PLANILLA_DEFAULT):
    """Dry-run: compara lo que importaría contra los totales col34 de la planilla.
    Distingue lo numérico (que la fórmula SUM sí cuenta) de lo capturado de celdas
    de TEXTO (que la planilla pierde). Cuadra si lo numérico coincide con col34."""
    wb = _abrir_planilla(ruta)
    ok = True
    total_texto = 0
    for hoja, mes in MESES.items():
        ws = wb[hoja]
        ndias = calendar.monthrange(ANIO, mes)[1]
        cols = range(COL_DIA1, COL_DIA1 + ndias)
        print(f"\n### {hoja} {ANIO}")
        etiquetas = {R_DELIVERY: "Pago al delivery", R_EFECTIVO: "Ventas Efectivo",
                     R_TRANSFER: "Ventas Transferencia", R_POS: "Ventas POS",
                     R_BOLETA: "Boleta Manual"}
        for r, et in etiquetas.items():
            num, txt = _split_num_txt([ws.cell(r, c).value for c in cols])
            total_texto += txt
            cache = num_celda(ws.cell(r, COL_MES).value)
            match = (num == cache)
            ok = ok and match
            extra = f"  (+{txt:,} de celdas de texto que la planilla no sumaba)" if txt else ""
            print(f"  [{'OK ' if match else 'DIF'}] {et:<22} numérico={num:>10,}  "
                  f"planilla(col34)={cache:>10,}{extra}")
        egr = sum(num_celda(ws.cell(fila, c).value) for fila in FILA_CATEGORIA for c in cols)
        egr_cache = sum(num_celda(ws.cell(fila, COL_MES).value) for fila in FILA_CATEGORIA)
        match = (egr == egr_cache)
        ok = ok and match
        print(f"  [{'OK ' if match else 'DIF'}] {'Egresos (todas cat)':<22} "
              f"importado={egr:>10,}  planilla(col34)={egr_cache:>10,}")
    if total_texto:
        print(f"\nNota: el importador captura ${total_texto:,} extra que la planilla "
              "tenía como texto (celdas 'deli 25500') y su fórmula SUM ignoraba.")
    print(("VALIDACIÓN OK: lo numérico reproduce la planilla exacto" if ok
           else "HAY DIFERENCIAS NUMÉRICAS: revisar antes de commitear"))
    return ok


def _limpiar_import_mes(fin, mes):
    desde, hasta = f"{ANIO}-{mes:02d}-01", f"{ANIO}-{mes:02d}-31"
    with fin._lock:
        fin._con.execute("DELETE FROM egresos WHERE origen='import' AND fecha BETWEEN ? AND ?",
                         (desde, hasta))
        fin._con.execute("DELETE FROM ingresos_manuales WHERE origen='import' AND fecha BETWEEN ? AND ?",
                         (desde, hasta))
        fin._con.execute("DELETE FROM ventas_dia WHERE origen='import' AND fecha BETWEEN ? AND ?",
                         (desde, hasta))
        fin._con.commit()


def importar(fin, ruta=PLANILLA_DEFAULT):
    """Carga MAYO/JUNIO/JULIO en el módulo de finanzas. Devuelve un resumen por mes."""
    wb = _abrir_planilla(ruta)
    resumen = {}
    for hoja, mes in MESES.items():
        ws = wb[hoja]
        _limpiar_import_mes(fin, mes)
        ndias = calendar.monthrange(ANIO, mes)[1]
        n_egr = n_vd = n_bm = 0
        for c in range(COL_DIA1, COL_DIA1 + ndias):
            dia = c - COL_DIA1 + 1
            fecha = f"{ANIO}-{mes:02d}-{dia:02d}"
            efe = num_celda(ws.cell(R_EFECTIVO, c).value)
            tr = num_celda(ws.cell(R_TRANSFER, c).value)
            pos = num_celda(ws.cell(R_POS, c).value)
            if efe or tr or pos:
                fin.set_ventas_dia(fecha, efe, tr, pos, origen="import"); n_vd += 1
            deli = num_celda(ws.cell(R_DELIVERY, c).value)
            if deli:
                fin.registrar_egreso(fecha, "Pago al delivery", deli,
                                     nota="import", origen="import"); n_egr += 1
            bm = num_celda(ws.cell(R_BOLETA, c).value)
            if bm:
                fin.registrar_ingreso_manual(fecha, bm, origen="import"); n_bm += 1
            for fila, catnom in FILA_CATEGORIA.items():
                m = num_celda(ws.cell(fila, c).value)
                if m:
                    prov = ws.cell(fila, 1).value
                    prov = str(prov).strip() if prov else None
                    nota = f"import; proveedor: {prov}" if prov else "import"
                    fin.registrar_egreso(fecha, catnom, m, nota=nota, origen="import")
                    n_egr += 1
        resumen[hoja] = {"dias_ventas": n_vd, "egresos": n_egr, "boletas": n_bm}
        print(f"  {hoja}: {n_vd} días de ventas, {n_egr} egresos, {n_bm} boletas manuales")
    return resumen


def main():
    ap = argparse.ArgumentParser(description="Importa el histórico de la planilla al libro de caja.")
    ap.add_argument("--planilla", default=str(PLANILLA_DEFAULT))
    ap.add_argument("--db", default=str(POS / "pedidos.db"))
    ap.add_argument("--dry-run", action="store_true", help="solo validar contra la planilla")
    ap.add_argument("--commit", action="store_true", help="escribir en la base")
    args = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    if args.dry_run or not args.commit:
        ok = validar(args.planilla)
        if not args.commit:
            print("\n(dry-run; usa --commit para escribir en la base)")
            sys.exit(0 if ok else 1)

    print(f"\nImportando a {args.db} ...")
    n = NucleoPedidos(db_path=Path(args.db))
    fin = NucleoFinanzas(n)
    importar(fin, args.planilla)
    # Verificación post-import: los totales del mes deben cuadrar con la planilla.
    print("\nVerificación post-import (resumen_mes vs planilla):")
    wb = _abrir_planilla(args.planilla)
    all_ok = True
    for hoja, mes in MESES.items():
        rm = fin.resumen_mes(ANIO, mes)
        ws = wb[hoja]
        e_ref = num_celda(ws.cell(R_EFECTIVO, COL_MES).value)
        t_ref = num_celda(ws.cell(R_TRANSFER, COL_MES).value)
        p_ref = num_celda(ws.cell(R_POS, COL_MES).value)
        match = (rm["efectivo"] == e_ref and rm["transferencia"] == t_ref and rm["pos"] == p_ref)
        all_ok = all_ok and match
        print(f"  [{'OK ' if match else 'DIF'}] {hoja}: efectivo {rm['efectivo']:,} / "
              f"transf {rm['transferencia']:,} / pos {rm['pos']:,}  (real={rm['ingreso_real']:,})")
    n._con.close()
    print("\n" + ("IMPORT OK" if all_ok else "IMPORT CON DIFERENCIAS"))
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()
