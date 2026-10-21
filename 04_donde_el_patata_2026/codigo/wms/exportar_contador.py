#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
exportar_contador.py — Consolidado mensual para el contador.
============================================================
Genera un .xlsx con el MISMO layout que la planilla de Christian (ingresos por
medio de pago + egresos por categoría + Real/Teórico + IVA neto), pero con UNA
columna = el total del mes, listo para entregarle al contador. El detalle diario
ya vive en el sistema; esto es el agregado que él necesita.

Reusa `NucleoFinanzas.resumen_mes` (la fuente única de verdad de los cálculos).

Uso:
  python wms/exportar_contador.py --anio 2026 --mes 5 --db <ruta.db> --salida cons.xlsx
Auto-test:
  python wms/exportar_contador.py --test
"""
import argparse
import sys
from datetime import datetime
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
POS = RAIZ / "pos"
sys.path.insert(0, str(POS))
sys.path.insert(0, str(AQUI))
from nucleo import NucleoPedidos           # noqa: E402
from finanzas import NucleoFinanzas         # noqa: E402

MES_NOMBRE = {1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL", 5: "MAYO", 6: "JUNIO",
              7: "JULIO", 8: "AGOSTO", 9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE",
              12: "DICIEMBRE"}


def exportar_mes(fin, anio, mes, salida):
    """Escribe el consolidado del mes en `salida` (.xlsx). Devuelve la ruta."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    rm = fin.resumen_mes(anio, mes)
    cats = fin.listar_categorias()  # en orden de la planilla
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{MES_NOMBRE[mes][:9]} {anio}"

    neg = Font(color="9C0006")
    bold = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    money = '#,##0;[Red]-#,##0'
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 16

    fila = [1]

    def escribir(a, b, valor=None, seccion=False, total=False, negrita=False):
        r = fila[0]
        ws.cell(r, 1, a)
        ws.cell(r, 2, b)
        if valor is not None:
            c = ws.cell(r, 3, int(valor))
            c.number_format = money
            if total or negrita:
                c.font = bold
            if valor < 0:
                c.font = neg
        if seccion:
            for col in (1, 2, 3):
                ws.cell(r, col).fill = hdr_fill
                ws.cell(r, col).font = hdr_font
        elif total:
            ws.cell(r, 1).font = bold
            ws.cell(r, 2).font = bold
        fila[0] += 1

    escribir(f"CONSOLIDADO {MES_NOMBRE[mes]} {anio}", "Donde el Patata — Huechuraba", seccion=True)
    escribir("", "")

    # --- INGRESOS ---
    escribir("INGRESOS", "", seccion=True)
    escribir("", "Ventas con Efectivo", rm["efectivo"])
    escribir("", "Ventas con Transferencia", rm["transferencia"])
    escribir("", "Ventas con POS", rm["pos"])
    escribir("", "Boleta Manual", rm["boleta_manual"])
    escribir("Depositar", "Boleta Efectiva (B.Manual − Transf.)", rm["boleta_efectiva"])
    escribir("Real", "Total Ingresos", rm["ingreso_real"], total=True)
    escribir("Teórico", "Total Ingresos", rm["ingreso_teorico"], total=True)
    escribir("", "")

    # --- EGRESOS (insumos variables + costos fijos, en orden de la planilla) ---
    escribir("EGRESOS", "", seccion=True)
    for cat in cats:
        if cat["tipo"] in ("insumo_variable", "costo_fijo"):
            monto = rm["por_categoria"].get(cat["nombre"], 0)
            escribir(cat.get("proveedor_sugerido") or "", cat["nombre"], monto)
    escribir("", "Total Egresos", rm["total_egresos"], total=True)
    escribir("", "")

    # --- Pago al delivery (egreso propio) ---
    escribir("", "Pago al delivery", rm["pago_delivery"])
    escribir("", "")

    # --- OTROS EGRESOS ---
    escribir("OTROS EGRESOS", "", seccion=True)
    for cat in cats:
        if cat["tipo"] == "otro_egreso":
            escribir("", cat["nombre"], rm["por_categoria"].get(cat["nombre"], 0))
    escribir("", "Total Otros Egresos", rm["otros_egresos"], total=True)
    escribir("", "")

    # --- FLUJO ---
    escribir("FLUJO DEL MES", "", seccion=True)
    escribir("Real", "Flujo (Ingreso Real − Egresos)", rm["flujo_real"], total=True)
    escribir("Teórico", "Flujo (Ingreso Teórico − Egresos)", rm["flujo_teorico"], total=True)
    escribir("", "")

    # --- TRIBUTARIO ---
    escribir("TRIBUTARIO", "", seccion=True)
    escribir("", "Compras con Factura", rm["compras_con_factura"])
    escribir("", "Compras sin Factura", rm["compras_sin_factura"])
    escribir("", "IVA Débito (ventas documentadas)", rm["iva_debito"])
    escribir("", "IVA Crédito (compras con factura)", rm["iva_credito"])
    escribir("", "IVA Neto (débito − crédito)", rm["iva_neto"], total=True)
    escribir("", "Caja no documentada (Real − Teórico)", rm["no_documentado"])

    escribir("", "")
    escribir("", f"Generado {datetime.now().strftime('%Y-%m-%d %H:%M')} · "
                 f"{rm['n_dias_con_datos']} día(s) con datos")

    salida = Path(salida)
    wb.save(salida)
    return salida


def _test():
    import tempfile
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    fails = [0]

    def chk(msg, cond):
        if not cond:
            fails[0] += 1
        print(f"  [{'OK' if cond else 'FALLA'}] {msg}")

    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        n = NucleoPedidos(db_path=Path(tmp) / "exp.db")
        fin = NucleoFinanzas(n)
        # Un mes chico y determinista.
        F = "2026-05-10"
        fin.set_ventas_dia(F, efectivo=100000, transferencia=30000, pos=50000)
        fin.registrar_ingreso_manual(F, 40000)
        fin.registrar_egreso(F, "Pan", 20000)
        fin.registrar_egreso(F, "Churrasco", 50000, con_factura=True)
        fin.registrar_egreso(F, "Luz", 30000)
        fin.registrar_egreso(F, "Reparaciones", 15000)
        fin.registrar_egreso(F, "Pago al delivery", 8000)
        salida = Path(tmp) / "cons_mayo.xlsx"
        exportar_mes(fin, 2026, 5, salida)
        chk("se generó el archivo .xlsx", salida.exists() and salida.stat().st_size > 0)

        import openpyxl
        ws = openpyxl.load_workbook(salida, data_only=True).active
        vals = {}
        for r in range(1, ws.max_row + 1):
            b = ws.cell(r, 2).value
            c = ws.cell(r, 3).value
            if b and c is not None:
                vals[str(b)] = c
        chk("Ventas con Efectivo = 100000", vals.get("Ventas con Efectivo") == 100000)
        chk("Boleta Manual = 40000", vals.get("Boleta Manual") == 40000)
        chk("Churrasco (egreso) = 50000", vals.get("Churrasco") == 50000)
        chk("Total Egresos = 100000", vals.get("Total Egresos") == 100000)
        chk("Pago al delivery = 8000", vals.get("Pago al delivery") == 8000)
        real = [ws.cell(r, 3).value for r in range(1, ws.max_row + 1)
                if ws.cell(r, 1).value == "Real" and "Total Ingresos" == ws.cell(r, 2).value]
        chk("Total Ingresos Real = 180000", real and real[0] == 180000)
        n._con.close()

    print("\n" + ("TODO OK" if fails[0] == 0 else f"{fails[0]} FALLARON"))
    return fails[0] == 0


def main():
    ap = argparse.ArgumentParser(description="Consolidado mensual para el contador (.xlsx).")
    ap.add_argument("--test", action="store_true")
    ap.add_argument("--anio", type=int, default=datetime.now().year)
    ap.add_argument("--mes", type=int, default=datetime.now().month)
    ap.add_argument("--db", default=str(POS / "pedidos.db"))
    ap.add_argument("--salida", default=None)
    args = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    if args.test:
        sys.exit(0 if _test() else 1)
    n = NucleoPedidos(db_path=Path(args.db))
    fin = NucleoFinanzas(n)
    salida = args.salida or f"consolidado_{args.anio}_{args.mes:02d}.xlsx"
    ruta = exportar_mes(fin, args.anio, args.mes, salida)
    n._con.close()
    print(f"Consolidado escrito en: {ruta}")


if __name__ == "__main__":
    main()
