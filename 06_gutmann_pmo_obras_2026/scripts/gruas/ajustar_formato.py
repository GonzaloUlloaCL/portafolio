"""Alinea las filas del presupuesto de Grúas (código, descripción, unidad y números), con el alto de cada fila
según su texto para que nada quede cortado.

Uso: python -I ajustar_formato.py <Presupuesto_Gruas.xlsx>
Escribe en el mismo archivo. No toca la cabecera (filas 1 a 6) ni los textos.
"""
import math
import sys

import openpyxl
from openpyxl.styles import Alignment

RUTA = sys.argv[1]
wb = openpyxl.load_workbook(RUTA)
ws = wb["Presupuesto"]
assert str(ws["A2"].value).startswith("PIEDRA ANGULAR"), "cabecera"


ANCHO_B, ANCHO_G = ws.column_dimensions["B"].width, ws.column_dimensions["G"].width
CAR = 1.15   # caracteres de Calibri 10 por unidad de ancho de columna, aproximado
for r in range(7, ws.max_row + 1):
    a, b, g = ws.cell(r, 1), ws.cell(r, 2), ws.cell(r, 7)
    es_partida = isinstance(a.value, str) and a.value.count(".") == 2 or b.value == "+ Agregar partida"
    if es_partida:
        a.alignment = Alignment(horizontal="left", vertical="top")
        b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(r, 3).alignment = Alignment(horizontal="center", vertical="top")
        for k in (4, 5, 6):
            ws.cell(r, k).alignment = Alignment(horizontal="right", vertical="top")
        g.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        lineas = max(1, math.ceil(len(str(b.value or "")) / (ANCHO_B * CAR)), math.ceil(len(str(g.value or "")) / (ANCHO_G * CAR)))
        ws.row_dimensions[r].height = 13.5 * lineas + 1.5
    elif isinstance(b.value, str) and b.value.startswith(("Subtotal", "TOTAL", "COSTO", "SUBTOTAL", "GASTOS", "UTILIDADES", "IVA")):
        ws.cell(r, 6).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(r, 5).alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[r].height = 16
for nombre, col in (("Léeme", "A"),):
    h = wb[nombre]
    ancho = h.column_dimensions[col].width
    for r in range(1, h.max_row + 1):
        texto = str(h[f"{col}{r}"].value or "")
        if len(texto) > ancho * CAR:
            h.row_dimensions[r].height = 15 * math.ceil(len(texto) / (ancho * CAR)) + 2
wb.save(RUTA)
print("formato ajustado:", RUTA)
