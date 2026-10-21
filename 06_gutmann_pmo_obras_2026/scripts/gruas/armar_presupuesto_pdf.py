"""PDF del presupuesto de reparaciones de Grúas Cavi para el cierre de la etapa (base: el de Talca): la hoja «Presupuesto» del Excel vigente, sin la columna
de observaciones y con los porcentajes de gastos generales, utilidades e IVA en blanco (los montos sí salen).

Uso: python -I armar_presupuesto_pdf.py <Presupuesto_Gruas.xlsx> <salida.pdf>
Cada subsector (A.1, B.1 … B.4) es una tabla enmarcada que nunca se corta entre páginas: lleva su encabezado de
columnas, sus partidas y su subtotal. El total de cada sector va pegado a su último subsector, y el resumen es un bloque.
Las cantidades son fórmulas del Excel sin valor guardado: se evalúan aquí, solo si son aritmética simple.
"""
import re, sys
from datetime import date
import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

excel, salida = sys.argv[1], sys.argv[2]
ws = openpyxl.load_workbook(excel)["Presupuesto"]
ARIT = re.compile(r"^[0-9.+\-*/() ]+$")
ORTOGRAFIA = []   # el impreso respeta cada texto del Excel tal cual (Gonza, 16-oct)
AZUL, CLARO, LINEA = colors.HexColor("#1F3A5F"), colors.HexColor("#E8EDF3"), colors.HexColor("#9AA8B8")


def numero(v):
    if v is None or isinstance(v, (int, float)):
        return v
    s = str(v).lstrip("=")
    if not ARIT.match(s):
        raise SystemExit(f"Cantidad no aritmética: {v!r}")
    return eval(s, {"__builtins__": {}})


def miles(v):
    return "$ " + f"{round(v):,}".replace(",", ".")


def cant(v):
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# recorrer la hoja: sectores con sus subsectores y partidas
sectores, actual, sub = [], None, None
encabezado = [ws.cell(r, 1).value for r in range(1, 7) if ws.cell(r, 1).value]
encabezado = ["FECHA: " + date.today().strftime("%d-%m-%Y") if str(t).startswith("FECHA") else t for t in encabezado]  # siempre la del día
for r in range(7, ws.max_row + 1):
    a, b, c = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
    if b == "RESUMEN":
        break
    if b == "+ Agregar partida":                                            # fila libre para Gonza: no se imprime
        continue
    if isinstance(b, str) and c == "Unidad":                               # A OBRAS PRELIMINARES …
        actual = {"titulo": b, "subs": []}; sectores.append(actual)
    elif isinstance(a, str) and re.match(r"^[A-E]\.\d+\.\d+$", a):         # partida
        q, pu = numero(ws.cell(r, 4).value), numero(ws.cell(r, 5).value)
        for x, y in ORTOGRAFIA:
            b = re.sub(x, y, b)
        sub["filas"].append((a, b, c or "", q, pu, (q or 0) * (pu or 0)))
    elif isinstance(b, str) and not b.upper().startswith(("SUBTOTAL", "SUB-TOTAL", "TOTAL")):   # A.1 Excavación …
        sub = {"titulo": b, "filas": []}; actual["subs"].append(sub)

cd = sum(f[5] for s in sectores for u in s["subs"] for f in u["filas"])
tasas = {ws.cell(r, 2).value: ws.cell(r, 5).value for r in range(1, ws.max_row + 1)}   # por rótulo: las filas se mueven
t_gg, t_ut, t_iva = (numero(tasas.get(k)) for k in ("GASTOS GENERALES", "UTILIDADES", "IVA"))   # del Excel; no se imprimen
assert all(isinstance(t, float) and 0 < t < 1 for t in (t_gg, t_ut, t_iva)), "faltan las tasas en el Excel"
gg, ut = cd * t_gg, cd * t_ut
neto = cd + gg + ut
iva = neto * t_iva

# estilos
p = lambda txt, **k: Paragraph(txt, ParagraphStyle("x", fontName=k.get("f", "Helvetica"), fontSize=k.get("s", 8.2),
                                                   leading=k.get("l", 10), alignment=k.get("a", 0),
                                                   textColor=k.get("c", colors.black)))
ANCHOS = [16 * mm, 88 * mm, 14 * mm, 20 * mm, 25 * mm, 28 * mm]   # 191 mm útiles en carta con márgenes de 12,5 mm
COLS = ["Ítem", "Descripción", "Unidad", "Cantidad", "P. Unitario", "Total"]


def tabla_sub(sector, s, primero, total_sector=None):
    filas = [[p(sector, f="Helvetica-Bold", s=9.5, c=colors.white)] + [""] * 5] if primero else []
    filas.append([p(s["titulo"], f="Helvetica-Bold", s=8.6)] + [""] * 5)
    filas.append([p(t, f="Helvetica-Bold", s=7.6, a=0 if i < 3 else 2) for i, t in enumerate(COLS)])
    for cod, desc, uni, q, pu, tot in s["filas"]:
        filas.append([p(cod), p(desc), p(uni, a=1), p(cant(q) if q is not None else "", a=2),
                      p(miles(pu) if pu is not None else "", a=2), p(miles(tot), a=2)])
    sub = sum(f[5] for f in s["filas"])
    filas.append(["", p("Subtotal " + s["titulo"].split(" ")[0], f="Helvetica-Bold"), "", "", "",
                  p(miles(sub), f="Helvetica-Bold", a=2)])
    if total_sector is not None:
        filas.append(["", p("TOTAL " + sector.split(" ")[0], f="Helvetica-Bold", s=9), "", "", "",
                      p(miles(total_sector), f="Helvetica-Bold", s=9, a=2)])
    i0 = 1 if primero else 0
    est = [("BOX", (0, 0), (-1, -1), 0.8, AZUL), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
           ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
           ("SPAN", (0, i0), (-1, i0)), ("BACKGROUND", (0, i0), (-1, i0), CLARO),
           ("LINEBELOW", (0, i0 + 1), (-1, i0 + 1), 0.6, AZUL),
           ("LINEBELOW", (0, i0 + 2), (-1, -2), 0.25, LINEA)]
    if primero:
        est += [("SPAN", (0, 0), (-1, 0)), ("BACKGROUND", (0, 0), (-1, 0), AZUL)]
    ultima_sub = -2 if total_sector is not None else -1
    est += [("LINEABOVE", (0, ultima_sub), (-1, ultima_sub), 0.6, AZUL), ("BACKGROUND", (0, ultima_sub), (-1, ultima_sub), CLARO)]
    if total_sector is not None:
        est += [("BACKGROUND", (0, -1), (-1, -1), AZUL), ("TEXTCOLOR", (0, -1), (-1, -1), colors.white)]
        filas[-1] = ["", p("TOTAL " + sector, f="Helvetica-Bold", s=9, c=colors.white), "", "", "",
                     p(miles(total_sector), f="Helvetica-Bold", s=9, a=2, c=colors.white)]
    t = Table(filas, colWidths=ANCHOS)
    t.setStyle(TableStyle(est))
    return KeepTogether([t, Spacer(1, 4 * mm)])


story = [p(encabezado[0], f="Helvetica-Bold", s=13, l=16), Spacer(1, 2 * mm)]
story += [p(t, s=9, l=11.5) for t in encabezado[1:]]
story.append(Spacer(1, 5 * mm))
for s in sectores:
    total = sum(f[5] for u in s["subs"] for f in u["filas"])
    for i, u in enumerate(s["subs"]):
        story.append(tabla_sub(s["titulo"], u, i == 0, total if i == len(s["subs"]) - 1 else None))

res = [[p("RESUMEN", f="Helvetica-Bold", s=9.5, c=colors.white), "", ""],
       [p("COSTO DIRECTO", f="Helvetica-Bold"), "", p(miles(cd), f="Helvetica-Bold", a=2)],
       [p("GASTOS GENERALES"), "", p(miles(gg), a=2)],
       [p("UTILIDADES"), "", p(miles(ut), a=2)],
       [p("SUBTOTAL NETO", f="Helvetica-Bold"), "", p(miles(neto), f="Helvetica-Bold", a=2)],
       [p("IVA"), "", p(miles(iva), a=2)],
       [p("TOTAL", f="Helvetica-Bold", s=10, c=colors.white), "", p(miles(neto + iva), f="Helvetica-Bold", s=10, a=2, c=colors.white)]]
rt = Table(res, colWidths=[118 * mm, 45 * mm, 28 * mm])   # la columna del medio es la del porcentaje: va en blanco
rt.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, AZUL), ("SPAN", (0, 0), (-1, 0)),
                        ("BACKGROUND", (0, 0), (-1, 0), AZUL), ("BACKGROUND", (0, -1), (-1, -1), AZUL),
                        ("LINEBELOW", (0, 1), (-1, -2), 0.25, LINEA), ("BACKGROUND", (0, 4), (-1, 4), CLARO),
                        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
story.append(KeepTogether([rt]))


def pie(canvas, doc):
    canvas.saveState(); canvas.setFont("Helvetica", 7.5); canvas.setFillColor(colors.HexColor("#5B6862"))
    canvas.drawString(12.5 * mm, 8 * mm, encabezado[3] if len(encabezado) > 3 else "")
    canvas.drawRightString(letter[0] - 12.5 * mm, 8 * mm, f"Página {doc.page}")
    canvas.restoreState()


SimpleDocTemplate(salida, pagesize=letter, leftMargin=12.5 * mm, rightMargin=12.5 * mm, topMargin=12 * mm,
                  bottomMargin=14 * mm, title="Presupuesto Grúas Cavi").build(story, onFirstPage=pie, onLaterPages=pie)
print(f"CD {miles(cd)} · neto {miles(neto)} · total {miles(neto + iva)} → {salida}")
