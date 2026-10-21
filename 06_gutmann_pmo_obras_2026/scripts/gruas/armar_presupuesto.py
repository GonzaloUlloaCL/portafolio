"""Arma el presupuesto de reparaciones de Grúas Cavi (Santa Alejandra 03520) con el formato de San Miguel y Talca.

Uso: python -I armar_presupuesto.py <salida.xlsx>
Trabajo menor sin plano: las partidas y cantidades salen de las 6 fotos de evidencias/ y de las notas de Gonza del
19-10-2026. Nombres del catálogo de San Miguel v2.0 donde existen; las tres partidas que no existen llevan el nombre
que Gonza les dio. El P.U. queda vacío (se pinta en rojo). Nunca escribe encima de un archivo existente: el vigente lo
edita Gonza.
"""
import os
import sys

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SALIDA = sys.argv[1]
if os.path.exists(SALIDA):
    sys.exit(f"{SALIDA} ya existe: el vigente lo edita Gonza. Genera a otro nombre.")

AZUL, CELESTE, AMARILLO, ROJO_CF = "FF1F3A5F", "FFE7ECF3", "FFFFF2CC", "FFF4B6B6"
FINO = Side(style="thin", color="FFBFBFBF")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)

# ------------------------------------------------------------------------------------------------- contenido
# Cada partida: (nombre, unidad, cantidad, observación). Cantidad: número o fórmula; None si falta.

ZONAS = [
    ("A", "SECTOR OFICINA", [
        ("Sector baño", [
            ("Reparación sanitaria (cambio de sifón de urinario)", "UN", 2,
             "Cambio de sifón (Gonza). Foto «Reparacion sanitaria»: 2 urinarios murales fuera de uso («NO USAR»)."),
        ]),
    ]),
    ("B", "SECTOR CAMARINES", [
        ("Sector duchas", [
            ("Habilitar y restaurar duchas", "UN", 2,
             "2 duchas (Gonza). Foto «Habilitar y restaurar duchas x2»: cerámica, perfiles de aluminio y desagüe de "
             "piso a la vista."),
        ]),
        ("Sector camarines", [
            ("Guardapolvo", "ml", 2.55,
             "Cambiar guardapolvo de 5 cm de alto: 2,55 ml (Gonza). Foto «Cambiar guardapolvo»: está despegado y roto "
             "junto al marco de la puerta."),
        ]),
        ("Sector caja escala", [
            ("Restauración y pintura muros", "m²", "=2.26*7.9",
             "1 muro: 2,26 de alto × 7,90 de largo (Gonza). Sin foto."),
            ("Restauración y pintura de cielo", "m²", "=7.9*1.1",
             "Cielo: 7,90 de largo × 1,10 de ancho (Gonza). Sin foto."),
        ]),
        ("Sector segundo piso", [
            ("Guardapolvo", "ml", 12.8,
             "Guardapolvo de 5 cm de alto: 12,80 ml (Gonza). Foto «Guardapolvo Segundo piso»: pasillo sin guardapolvo."),
            ("Restauración chapa comedor (puerta metálica)", "UN", 1,
             "Chapa sobrepuesta de la puerta metálica del comedor (Gonza). Fotos «Chapa comedor Segundo piso» 1 y 2: "
             "las dos caras de la misma puerta, chapa con óxido."),
        ]),
    ]),
]

# ------------------------------------------------------------------------------------------------- libro
wb = openpyxl.Workbook()
leeme = wb.active
leeme.title = "Léeme"
ws = wb.create_sheet("Presupuesto")


def fuente(**kw):
    return Font(name="Calibri", sz=kw.pop("sz", 10), **kw)


# Cabecera fija: filas 1 a 6 combinadas A:F, Calibri 10 centrado, negrita de la 2 a la 6, color automático
CABECERA = ["PRESUPUESTO DESGLOSADO POR ITEMS A SUMA ALZADA", "PIEDRA ANGULAR SPA 77.453.608-6",
            "LAS SIERVAS DE MARIA 5556, HUECHURABA", "OBRA: REPARACIONES GRÚAS CAVI",
            "DIRECCIÓN: Santa Alejandra 03520", "FECHA: 19-10-2026"]
for i, texto in enumerate(CABECERA, start=1):
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    c = ws.cell(i, 1, texto)
    c.font = fuente(b=i > 1)
    c.alignment = Alignment(horizontal="center")
if not str(ws["A2"].value).startswith("PIEDRA ANGULAR"):
    sys.exit("La cabecera no empieza con PIEDRA ANGULAR en A2.")

for col, ancho in zip("ABCDEFG", (8.9, 60.0, 9.0, 12.0, 14.0, 16.0, 90.0)):
    ws.column_dimensions[col].width = ancho


def fila_zona(r, letra, titulo):
    for k, v in enumerate(("", f"{letra} {titulo}", "Unidad", "Cantidad", "P. Unitario", "Total", "Observación"), 1):
        c = ws.cell(r, k, v or None)
        c.font = fuente(b=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="left" if k == 2 else "center")


def fila_celeste(r, texto, total=None, color_texto=None):
    for k in range(1, 8):
        c = ws.cell(r, k)
        c.fill = PatternFill("solid", fgColor=CELESTE)
        c.font = fuente()
    ws.cell(r, 2, texto).font = fuente(b=True, color=color_texto)
    if total:
        c = ws.cell(r, 6, total)
        c.font = fuente(b=True)
        c.number_format = "#,##0"


def fila_partida(r, codigo, nombre, unidad, obs, cantidad=None):
    valores = (codigo, nombre, unidad, cantidad, None, f"=D{r}*E{r}", obs)
    for k, v in enumerate(valores, 1):
        c = ws.cell(r, k, v)
        c.font = fuente()
        c.border = BORDE
        c.alignment = Alignment(vertical="top", wrap_text=k in (2, 7))
    for k, formato in ((4, "#,##0.00"), (5, "#,##0"), (6, "#,##0")):
        ws.cell(r, k).number_format = formato
    for k in (4, 5):
        ws.cell(r, k).fill = PatternFill("solid", fgColor=AMARILLO)


def fila_subtotal(r, texto, desde, hasta):
    for k in range(1, 8):
        c = ws.cell(r, k)
        c.border = BORDE
        c.font = fuente()
    c = ws.cell(r, 2, texto)
    c.font = fuente(b=True)
    c.alignment = Alignment(horizontal="right")
    c = ws.cell(r, 6, f"=SUM(F{desde}:F{hasta})")
    c.font = fuente(b=True)
    c.number_format = "#,##0"


r = 7
totales_zona = []
for letra, titulo, sectores in ZONAS:
    fila_zona(r, letra, titulo)
    r += 1
    subtotales = []
    for s, (nombre_sector, partidas) in enumerate(sectores, 1):
        cod = f"{letra}.{s}"
        fila_celeste(r, f"{cod} {nombre_sector}", color_texto=AZUL)
        r += 1
        desde = r
        for p, (nombre, unidad, cantidad, obs) in enumerate(partidas, 1):
            fila_partida(r, f"{cod}.{p}", nombre, unidad, obs, cantidad)
            r += 1
        if s == len(sectores):  # fila libre al final de cada zona, dentro de su subtotal
            fila_partida(r, None, "+ Agregar partida", None, None)
            r += 1
        fila_subtotal(r, f"Subtotal {cod}", desde, r - 1)
        subtotales.append(f"F{r}")
        r += 1
    fila_celeste(r, f"TOTAL {letra}", "=" + "+".join(subtotales))
    totales_zona.append(f"F{r}")
    r += 2

fila_celeste(r, "RESUMEN")
r += 1
fila_celeste(r, "COSTO DIRECTO", "=" + "+".join(totales_zona))
f_cd = r
r += 2
filas_resumen = {}
for nombre, tasa in (("GASTOS GENERALES", 0.05), ("UTILIDADES", 0.15)):
    ws.cell(r, 2, nombre).font = fuente()
    c = ws.cell(r, 5, tasa)
    c.font = fuente(color="FF000000")
    c.fill = PatternFill("solid", fgColor=AMARILLO)
    c.number_format = "0%"
    c = ws.cell(r, 6, f"=F{f_cd}*E{r}")
    c.font = fuente()
    c.number_format = "#,##0"
    filas_resumen[nombre] = r
    r += 2
ws.cell(r, 2, "SUBTOTAL NETO").font = fuente()
c = ws.cell(r, 6, f"=F{f_cd}+F{filas_resumen['GASTOS GENERALES']}+F{filas_resumen['UTILIDADES']}")
c.font = fuente()
c.number_format = "#,##0"
f_neto = r
r += 2
ws.cell(r, 2, "IVA").font = fuente()
c = ws.cell(r, 5, 0.19)
c.font = fuente(color="FF000000")
c.fill = PatternFill("solid", fgColor=AMARILLO)
c.number_format = "0%"
c = ws.cell(r, 6, f"=F{f_neto}*E{r}")
c.font = fuente()
c.number_format = "#,##0"
f_iva = r
r += 2
ws.cell(r, 2, "TOTAL").font = fuente(b=True)
c = ws.cell(r, 6, f"=F{f_neto}+F{f_iva}")
c.font = fuente(b=True)
c.number_format = "#,##0"
ultima = r

roja = PatternFill("solid", fgColor=ROJO_CF, bgColor=ROJO_CF)
ws.conditional_formatting.add(f"D8:D{ultima}", FormulaRule(formula=["AND(LEN($A8)>4,ISBLANK($D8))"], fill=roja))
ws.conditional_formatting.add(f"E8:E{ultima}", FormulaRule(formula=["AND(LEN($A8)>4,ISBLANK($E8))"], fill=roja))
ws.freeze_panes = "A8"

# Léeme
leeme.column_dimensions["A"].width = 130
LEEME = [
    ("t", "PRESUPUESTO DESGLOSADO POR ÍTEMS A SUMA ALZADA — BORRADOR PARA REVISIÓN"),
    ("p", "Obra: reparaciones en Grúas Cavi, Santa Alejandra 03520. Trabajo menor sin plano: sector Oficina (baño) y "
          "sector Camarines (duchas, guardapolvos, caja escala y chapa del comedor)."),
    ("p", "Preparado para Luciano por Gonza el 19-10-2026, con la forma de San Miguel."),
    ("", ""),
    ("s", "Cómo está ordenado"),
    ("p", "1. Hoja «Presupuesto»: A Sector oficina y B Sector camarines. Al final, el resumen con GG 5 %, utilidades "
          "15 % e IVA."),
    ("p", "2. Las cantidades son las medidas de Gonza en terreno; el P.U. queda vacío y se pinta en rojo hasta que se "
          "complete. La «Observación» dice de qué foto y de qué nota sale cada partida."),
    ("p", "3. Cada zona termina con una fila «+ Agregar partida» en blanco, ya dentro de su subtotal."),
    ("p", "4. Fotos de respaldo en la carpeta evidencias/ de la obra."),
]
for i, (tipo, texto) in enumerate(LEEME, 1):
    c = leeme.cell(i, 1, texto or None)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if tipo == "t":
        c.font = Font(name="Calibri", sz=13, b=True, color=AZUL)
    elif tipo == "s":
        c.font = Font(name="Calibri", sz=11, b=True, color=AZUL)
    else:
        c.font = Font(name="Calibri", sz=11, color="FF000000")

wb.active = wb.sheetnames.index("Presupuesto")
wb.save(SALIDA)
print(f"{SALIDA}: filas del presupuesto 7 a {ultima}")
