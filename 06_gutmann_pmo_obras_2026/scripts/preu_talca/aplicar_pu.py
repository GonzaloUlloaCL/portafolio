"""Carga en el presupuesto de Talca los P.U. de las obras anteriores de Luciano, partida por partida y por nombre.

Uso: python -I aplicar_pu.py <presupuesto_talca.xlsx>
Lee la columna E (P. Unitario) de los presupuestos vigentes en el Drive: San Miguel (catálogo), San Bernardo y el Dren.
Busca cada partida de Talca por su nombre normalizado (sin tildes ni mayúsculas); primero con el nombre completo y,
si no aparece, sin el «(h = 3,00 m)». Prioridad: San Miguel, luego San Bernardo, luego el Dren. Si una obra trae el
mismo nombre con precios distintos, usa el más repetido (en empate, el de su zona D, las salas nuevas) e informa las
diferencias. Escribe solo la columna E de las filas con código y E vacía, con un comentario de la fuente. No toca
cantidades, cabecera ni las filas «+ Agregar partida».
"""
import re
import sys
import unicodedata
from collections import Counter

import openpyxl
from openpyxl.comments import Comment

DRIVE = "/mnt/chromeos/shared/GoogleDrive/MyDrive/Luciano_Obra"
OBRAS = [("San Miguel", f"{DRIVE}/San Miguel/Presupuesto_San_Miguel.xlsx"),
         ("San Bernardo", f"{DRIVE}/San_Bernardo/Presupuesto_San_Bernardo.xlsx"),
         ("Dren NEWAIR", f"{DRIVE}/Dren_Frances_(6x4x3)/Presupuesto_Dren_Frances.xlsx")]
CODIGO = re.compile(r"^[A-Z]\.\d+\.\d+$")
ARIT = re.compile(r"^[0-9.+\-*/() ]+$")
REF = re.compile(r"^=?F(\d+)/D(\d+)$")


def normal(s, sin_alto=False):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    if sin_alto:
        s = re.sub(r"\(h\s*=\s*[\d,.]+\s*m\)", "", s)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def numero(v, ws):
    """P.U. como número: valor, aritmética simple o F/D de la misma fila. None si no se puede."""
    if v is None or isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    m = REF.match(s)
    if m:
        f, d = numero(ws[f"F{m[1]}"].value, ws), numero(ws[f"D{m[2]}"].value, ws)
        return f / d if f is not None and d else None
    s = s.lstrip("=")
    return eval(s, {"__builtins__": {}}) if ARIT.match(s) else None


def catalogo():
    filas, malos = [], []
    for obra, ruta in OBRAS:
        ws = openpyxl.load_workbook(ruta)["Presupuesto"]
        for r in ws.iter_rows(min_row=7):
            cod = r[0].value
            if not (isinstance(cod, str) and CODIGO.match(cod.strip()) and r[1].value):
                continue
            pu = numero(r[4].value, ws)
            if pu is None:
                malos.append(f"{obra} {cod}: P.U. «{r[4].value}» no se pudo leer")
                continue
            filas.append(dict(obra=obra, cod=cod.strip(), nombre=str(r[1].value).strip(), pu=round(pu)))
    return filas, malos


def elegir(cands):
    """Entre las filas de una misma obra: el P.U. más repetido; en empate, el de la zona D."""
    cuenta = Counter(c["pu"] for c in cands)
    mejor = max(cuenta.values())
    top = [c for c in cands if cuenta[c["pu"]] == mejor]
    return next((c for c in top if c["cod"].startswith("D.")), top[0])


def buscar(nombre, filas):
    for sin_alto in (False, True):
        clave = normal(nombre, sin_alto)
        for obra, _ in OBRAS:
            cands = [f for f in filas if f["obra"] == obra and normal(f["nombre"], sin_alto) == clave]
            if cands:
                otros = sorted({(f["obra"], f["pu"]) for f in filas if normal(f["nombre"], sin_alto) == clave})
                return elegir(cands), otros
    return None, []


def pesos(v):
    return f"{v:,.0f}".replace(",", ".")


def main(ruta):
    filas, malos = catalogo()
    wb = openpyxl.load_workbook(ruta)
    ws = wb["Presupuesto"]
    puestos, sin, difs = [], [], {}
    for r in ws.iter_rows(min_row=8):
        cod, nombre, e = r[0].value, r[1].value, r[4]
        if not (isinstance(cod, str) and CODIGO.match(cod)) or e.value is not None:
            continue
        f, otros = buscar(nombre, filas)
        if not f:
            sin.append(f"{cod} {nombre}")
            continue
        e.value = f["pu"]
        e.comment = Comment(f"{f['obra']} {f['cod']}, {pesos(f['pu'])}", "Alma")
        puestos.append((cod, nombre, f))
        if len({pu for _, pu in otros}) > 1:
            difs[normal(nombre)] = (nombre, otros, f["pu"])
    a2, a3 = str(ws["A2"].value), str(ws["A3"].value)
    if not a2.startswith("PIEDRA ANGULAR") or a3 != "LAS SIERVAS DE MARIA 5556, HUECHURABA":
        sys.exit("La cabecera no es la oficial: no se guarda.")
    wb.save(ruta)
    print(f"{len(puestos)} partidas con P.U.")
    for o in sorted({f['obra'] for *_, f in puestos}):
        print(f"  {o}: {sum(1 for *_, f in puestos if f['obra'] == o)}")
    print("Sin P.U.:", *sin or ["ninguna"], sep="\n  ")
    print("Diferencias entre obras (se usó el valor indicado):")
    for nombre, otros, pu in difs.values():
        print(f"  {nombre}: usado {pesos(pu)} · " + ", ".join(f"{o} {pesos(v)}" for o, v in otros))
    for m in malos:
        print("  No leído:", m)


if __name__ == "__main__":
    main(sys.argv[1])
