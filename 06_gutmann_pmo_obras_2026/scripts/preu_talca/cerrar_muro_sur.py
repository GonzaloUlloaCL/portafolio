"""El vano bajo la fachada sur no es acceso: es muro (Gonza, 17-10-2026). El hall suma esa cara a su guardapolvo y a la
pintura de muros (de 4,36 a 7,21 ml: de la Sala 06 a la Sala 05) y la cabecera toma la fecha del día.

Uso: python -I cerrar_muro_sur.py <Presupuesto_Talca.xlsx> <fecha dd-mm-aaaa>
"""
import sys

import openpyxl

RUTA, FECHA = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(RUTA)
ws = wb["Presupuesto"]
assert str(ws["A2"].value).startswith("PIEDRA ANGULAR"), "cabecera"
ws["A6"].value = "FECHA: " + FECHA
for f in ws.iter_rows(min_row=8):
    cod = f[0].value
    if cod in ("D.1.3", "D.1.5"):
        assert "4.36" in f[3].value, cod
        f[3].value = f[3].value.replace("4.36", "7.21")
        f[6].value = f[6].value.replace("muro sur 4,36 = 70,79 ml", "muro sur 7,21 (de la Sala 06 a la Sala 05: es muro, no acceso) = 73,64 ml").replace("4,36", "7,21")
wb.save(RUTA)
print("muro sur cerrado en", RUTA)
