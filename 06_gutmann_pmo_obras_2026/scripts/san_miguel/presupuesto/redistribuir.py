import sys, re, copy, openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Alignment
ENT, SAL = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(ENT); ws = wb["Presupuesto"]
SV = "Sin descontar puertas ni ventanas (criterio acordado con el mandante)."
TE = "medido en terreno el 08-10-2026"

# ---- datos actuales por código
IT = {}
for r in range(1, 400):
    c = ws.cell(r, 1).value
    if isinstance(c, str) and re.match(r"^[A-C]\.\d+\.\d+\.\d+$", c):
        IT[c] = [ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 7).value]
# ---- estilos de filas tipo
def estilo(r): return [copy.copy(ws.cell(r, k)._style) for k in range(1, 8)], ws.row_dimensions[r].height
ST = {"hdr": estilo(7), "sub": estilo(20), "item": estilo(21), "tot": estilo(17), "vacia": estilo(364),
      "res": estilo(365), "cd": estilo(366), "pct": estilo(368), "fin": estilo(376)}
RES = [(ws.cell(r, 2).value, ws.cell(r, 5).value, ws.cell(r, 7).value) for r in (366, 368, 370, 372, 374, 376)]

def it(cod, desc=None, uni=None, cant=None, obs=None):
    d, u, q, o = IT[cod]
    return ["it", desc or d, uni or u, cant if cant is not None else q, obs or o, cod]
def nuevo(desc, uni, cant, obs): return ["it", desc, uni, cant, obs, None]
def rell(area_txt, area_f, nom):
    return [nuevo("Relleno de piso", "m³", f"=({area_f})*0.7", f"{nom}: {area_txt} m² × 0,70 m de espesor. El piso de la casona se retira y se rellena."),
            nuevo("Losa radier", "m²", f"={area_f}", f"Radier sobre el relleno: {area_txt} m² ({nom}).")]

CAS1 = "Casona piso 1"
SEC = [
 ("A", "OBRAS PRELIMINARES", [
   ("Instalación de faenas y trabajos previos", [it(f"A.1.1.{i}") for i in range(1, 9)])]),
 ("B", "CASONA PISO 1", [
   ("Sector recepcion", [it("A.2.1.1"), it("A.2.1.2"), it("A.2.1.3"), it("A.2.1.4"),
      *rell("8,94", "13.08-4.14", "Recepción sin el bow window (13,08 rotulados − 4,14)"),
      nuevo("Pavimento (material por confirmar)", "m²", "=13.08-4.14", "Superficie útil rotulada de la Recepción (13,08 m²) menos el bow window (4,14 m², va en su sector) = 8,94 m². Material por definir."),
      nuevo("Zócalo / guardapolvo", "ml", 8.1, "Muros de la Recepción a escala, sin el bow window: fachada 0,86 + 0,69, muro de la chimenea 4,65 y costados 2 × 0,40, machones 0,30 + 0,80 = 8,10 ml. " + SV),
      it("B.4.1.6"), it("B.4.1.7")]),
   ("Sector arco ventana", [it("A.2.2.1"), it("A.2.2.2"), it("A.2.2.3"), it("A.2.2.4"), it("A.2.2.5"), it("A.2.2.6"),
      *rell("4,14", "4.14", "Bow window"),
      nuevo("Pavimento (material por confirmar)", "m²", 4.14, "Superficie interior del bow window, a escala: 4,14 m². Material por definir."),
      nuevo("Zócalo / guardapolvo", "ml", 5.42, "Contorno interior del bow window, a escala: 5,42 ml. " + SV)]),
   ("Sector sala de estar", [it("A.2.3.1"), it("A.2.3.2"), it("A.2.3.3"), it("A.2.3.4"), it("A.2.3.5"),
      it("B.2.2.2"), it("B.2.2.3"), *rell("22,96", "22.96", "Estar de Alumnos"),
      it("B.4.2.1"), it("B.4.2.2"), it("B.4.2.4"), it("B.4.2.6"), it("B.4.2.7")]),
   ("Sector salida posterior", [it("A.2.4.1"), it("A.2.4.2"), it("A.2.4.3"), it("A.2.4.4"), it("A.2.4.5"), it("A.2.4.6"),
      *rell("4,57", "4.57", "Salida posterior (2,75 × 1,66)"),
      it("B.4.3.1"), it("B.4.3.2"), it("B.4.3.4"), it("B.4.3.6"), it("B.4.3.7")]),
   ("Sector oficina 1", [it("A.2.5.1"), it("A.2.5.2"), it("A.2.5.3"), it("A.2.5.4"), it("A.2.5.5"), it("A.2.5.6"),
      nuevo("Tabiquería interior", "m²", "=(0.75+1.05+1.3)*3.5", "Muros nuevos (rosados) de la Oficina 01, a escala: 0,75 + 1,05 + 1,30 = 3,10 ml × 3,50 m = 10,85 m². Material por definir."),
      *rell("10,99", "10.99", "Oficina 01"),
      it("B.4.4.1"), it("B.4.4.2"), it("B.4.4.4"), it("B.4.4.6"), it("B.4.4.7")]),
   ("Sector B hombres 2", [it("A.2.6.1"), it("A.2.6.2"), it("A.2.6.3"), it("A.2.6.4"), it("A.2.6.5"),
      nuevo("Tabiquería interior", "m²", "=(1.65+2.05)*3.5", "Muros nuevos (rosados) del B. Hombres 2, a escala: 1,65 + 2,05 = 3,70 ml × 3,50 m = 12,95 m². Material por definir."),
      *rell("10,75", "5.35*2.05-0.25*0.85", "B. Hombres 2"),
      it("B.6.5.1"), it("B.6.5.2"), it("B.6.5.4"), it("B.6.5.5"), it("B.6.5.6"), it("B.6.5.7"), it("B.6.5.8")]),
   ("Sector B h+m 1", [it("A.2.7.1"), it("A.2.7.2"), it("A.2.7.3"), it("A.2.7.4"), it("A.2.7.5"), it("A.2.7.6"),
      *rell("5,84", "1.47*2.02+1.42*2.02", "B. Mujeres 1 (2,97) + B. Hombres 1 (2,87)"),
      it("B.6.1.1", desc="Pavimento pasillo de los baños (porcelanato o cerámica por confirmar)"),
      it("B.6.1.2", desc="Revestimiento de muros del pasillo de los baños"),
      it("B.6.1.3", desc="Cielo y pintura del pasillo de los baños"), it("B.6.1.4", desc="Puerta del pasillo de los baños"),
      *[it(f"B.6.2.{i}", desc=IT[f"B.6.2.{i}"][0] + " — B. Mujeres 1") for i in range(1, 9)],
      *[it(f"B.6.3.{i}", desc=IT[f"B.6.3.{i}"][0] + " — B. Hombres 1") for i in range(1, 9)]]),
   ("Sector hall y circulación", [
      *rell("27,98", "27.98", "Hall y circulación de la casona, a escala: hall 15,12 + resto de la pieza del frente 9,25 + pasillo de los baños 1 3,61")]),
 ]),
 ("C", "CASONA PISO 2", [
   ("Sector pieza 1", [*[it(f"A.3.1.{i}") for i in range(1, 5)],
      nuevo("Terciado", "m²", "=5.1*4.1", "Pieza 1: 5,10 × 4,10 = 20,91 m². Medidas interiores tomadas en terreno el 08-10-2026 (croquis del piso 2)."),
      nuevo("Piso flotante", "m²", "=5.1*4.1", "Igual que el terciado: 20,91 m²."),
      *[it(f"B.5.2.{i}") for i in range(1, 7)]]),
   ("Sector pieza 2", [*[it(f"A.3.2.{i}") for i in range(1, 5)],
      nuevo("Terciado", "m²", "=5.1*3.3", "Pieza 2: 5,10 × 3,30 = 16,83 m². Medidas interiores tomadas en terreno el 08-10-2026 (croquis del piso 2)."),
      nuevo("Piso flotante", "m²", "=5.1*3.3", "Igual que el terciado: 16,83 m²."),
      it("B.5.3.1"), it("B.5.3.2"), it("B.5.3.4"), it("B.5.3.5"), it("B.5.3.6")]),
   ("Sector pieza 3", [*[it(f"A.3.3.{i}") for i in range(1, 4)],
      nuevo("Terciado", "m²", "=1.5*4.1", "Pieza 3: 1,50 × 4,10 = 6,15 m². Medidas interiores tomadas en terreno el 08-10-2026 (croquis del piso 2)."),
      nuevo("Piso flotante", "m²", "=1.5*4.1", "Igual que el terciado: 6,15 m²."),
      *[it(f"B.5.4.{i}") for i in range(1, 7)]]),
   ("Sector pieza 4", [it("A.3.4.1"), it("A.3.4.2"),
      nuevo("Terciado", "m²", "=3.3*4.1", "Pieza 4: 3,30 × 4,10 = 13,53 m². Medidas interiores tomadas en terreno el 08-10-2026 (croquis del piso 2)."),
      nuevo("Piso flotante", "m²", "=3.3*4.1", "Igual que el terciado: 13,53 m²."),
      *[it(f"B.5.5.{i}") for i in range(1, 7)]]),
   ("Sector pieza 5", [it("A.3.5.1"), it("A.3.5.2", uni="m²"),
      nuevo("Terciado", "m²", "=3*2.5+2*2.4", "Pieza 5: 3,00 × 2,50 + 2,00 × 2,40 = 12,30 m². Medidas interiores tomadas en terreno el 08-10-2026 (croquis del piso 2)."),
      nuevo("Piso flotante", "m²", "=3*2.5+2*2.4", "Igual que el terciado: 12,30 m²."),
      *[it(f"B.5.6.{i}") for i in range(1, 7)]]),
   ("Sector baño", [*[it(f"A.3.6.{i}") for i in range(1, 6)], *[it(f"B.5.7.{i}") for i in range(1, 8)]]),
   ("Sector pasillo", [it("A.3.7.1"), it("A.3.7.2"),
      nuevo("Terciado", "m²", "=2*1.4+1*3.6", "Pasillo en L: 2,00 × 1,40 + 1,00 × 3,60 = 6,40 m², sin el hueco de la escalera. Medidas interiores tomadas en terreno el 08-10-2026 (croquis del piso 2)."),
      nuevo("Piso flotante", "m²", "=2*1.4+1*3.6", "Igual que el terciado: 6,40 m²."),
      *[it(f"B.5.1.{i}") for i in range(1, 7)]]),
 ]),
 ("D", "SALAS NUEVAS", [
   (f"Sector sala de capacitación {s:02d}", [it(f"B.3.{s}.9"),
      nuevo("Techumbre", "m²", f"='Cubicación base'!$B${s + 3}", f"Superficie útil rotulada de la Sala {s:02d}. La cubierta real suma espesor de muros y pendiente: falta el corte de techumbre."),
      *[it(f"B.3.{s}.{i}") for i in range(1, 9)]]) for s in range(1, 8)]),
 ("E", "BAÑOS NUEVOS", [
   ("Sector B. Mujeres 2", [it("B.6.4.9"), *[it(f"B.6.4.{i}") for i in range(1, 9)]]),
   ("Sector B. Universal", [it("B.6.6.9"), *[it(f"B.6.6.{i}") for i in range(1, 9)]]),
 ]),
 ("F", "GENERALES", [
   ("Estructura", [it("B.2.2.1")]),
   ("Techumbre y fachada", [it("B.7.1.2"), it("B.7.1.3"), it("B.7.2.1"), it("B.7.2.2"), it("B.7.2.3")]),
   ("Cierres exteriores", [it(f"C.1.1.{i}") for i in range(1, 5)]),
   ("Instalaciones generales", [it(f"C.2.1.{i}") for i in range(1, 5)]),
   ("Aseo", [it("C.3.1.1"), it("C.3.1.2")]),
 ]),
]
# ---- control: cada código viejo usado una vez o quitado a propósito
QUITADOS = {"B.4.1.1", "B.4.1.2", "B.4.1.3", "B.4.1.4", "B.4.1.5", "B.4.2.3", "B.4.2.5", "B.4.3.3", "B.4.3.5", "B.4.4.3", "B.4.4.5",
            "B.6.5.3", "B.5.3.3", "B.1.1.1", "B.1.1.2", "B.2.1.1", "B.2.1.2", "B.2.2.4", "B.7.1.1"}
usados = [x[5] for _, _, sec in SEC for _, its in sec for x in its if x[5]]
from collections import Counter
dup = [k for k, v in Counter(usados).items() if v > 1]
falt = sorted(set(IT) - set(usados) - QUITADOS)
assert not dup and not falt, (dup, falt)

# ---- mapa de códigos viejos -> nuevos, para corregir referencias en observaciones
MAPA = {}
for L, _, sec in SEC:
    for j, (_, its) in enumerate(sec, 1):
        for k, x in enumerate(its, 1):
            if x[5]: MAPA[x[5]] = f"{L}.{j}.{k}"

# ---- borrar desde la fila 7 y reescribir
ws.delete_rows(7, ws.max_row)
ws.conditional_formatting = type(ws.conditional_formatting)()
def poner(r, tipo, vals):
    sty, h = ST[tipo]
    for k in range(7):
        c = ws.cell(r, k + 1); c._style = copy.copy(sty[k]); c.value = vals[k] if k < len(vals) else None
    ws.row_dimensions[r].height = h
def remap(t):
    if not isinstance(t, str): return t
    return re.sub(r"\b([A-C]\.\d+\.\d+\.\d+)\b", lambda m: MAPA.get(m[1], m[1]), t)
r = 7; totales = []; primera = None
for L, nom, sec in SEC:
    poner(r, "hdr", [f"{L}. {nom}", f"{L} {nom}", "Unidad", "Cantidad", "P. Unitario", "Total", "Observación"]); r += 1
    ini = r
    for j, (snom, its) in enumerate(sec, 1):
        poner(r, "sub", [None, f"{L}.{j} {snom}"]); r += 1
        for k, x in enumerate(its, 1):
            _, d, u, q, o, old = x
            if isinstance(q, str) and re.fullmatch(r"=D\d+", q):   # referencia a otra fila: se reemplaza por su valor de origen
                q = IT["B.1.1.1"][2]
            poner(r, "item", [f"{L}.{j}.{k}", d, u, q, None, f"=D{r}*E{r}", remap(o)])
            ws.cell(r, 4).number_format = "#,##0.00"
            ws.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")
            primera = primera or r
            r += 1
    poner(r, "tot", [None, f"TOTAL {L}", None, None, None, f"=SUM(F{ini}:F{r - 1})"]); totales.append(r); r += 2
ult_item = r - 3
poner(r, "res", [None, "RESUMEN"]); r += 1
cd = r
for i, (nom, pct, obs) in enumerate(RES):
    tipo = "cd" if i == 0 else ("fin" if nom == "TOTAL" else "pct")
    if i == 0: f = "=" + "+".join(f"F{t}" for t in totales)
    elif nom == "GASTOS GENERALES" or nom == "UTILIDADES": f = f"=F{cd}*E{r}"
    elif nom == "SUBTOTAL NETO": f = f"=F{cd}+F{cd + 2}+F{cd + 4}"
    elif nom == "IVA": f = f"=F{cd + 6}*E{r}"
    else: f = f"=F{cd + 6}+F{cd + 8}"
    poner(r, tipo, [None, nom, None, None, pct, f, obs]); r += 2
ws.conditional_formatting.add(f"D{primera}:D{ult_item}", FormulaRule(formula=[f'AND(LEN($A{primera})>4,ISBLANK($D{primera}))'],
                              fill=PatternFill(start_color="F4B6B6", end_color="F4B6B6", fill_type="solid")))

# ---- Pendientes y Léeme: códigos nuevos
pe = wb["Pendientes"]
for row in pe.iter_rows(min_row=4):
    for c in row:
        if isinstance(c.value, str): c.value = remap(c.value)
pe["D4"].value = "F.2 Techumbre y fachada; ventanas, cielos y pavimentos"
le = wb["Léeme"]
for row in le.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("1. Hoja «Presupuesto»"):
            c.value = ("1. Hoja «Presupuesto», por zonas y sectores: A Obras preliminares, B Casona piso 1, C Casona piso 2, "
                       "D Salas nuevas, E Baños nuevos y F Generales. Cada sector (recinto) aparece una sola vez, con sus tareas en orden: "
                       "trabajos previos, obra gruesa y terminaciones.")
wb.save(SAL)
print("filas", r, "| totales", totales, "| CD fila", cd)
import json; json.dump(MAPA, open(SAL + ".mapa.json", "w"))
