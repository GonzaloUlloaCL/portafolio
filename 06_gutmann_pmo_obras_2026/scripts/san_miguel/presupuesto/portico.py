import copy, re, sys, openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange
src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src); ws = wb["Presupuesto"]
REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "Presupuesto" in c.value and c.value.startswith("="):
                sys.exit(f"referencia externa a Presupuesto en {s.title}!{c.coordinate}")
assert ws.merged_cells.ranges == set() or print("merged", ws.merged_cells.ranges)

def fila(t, col=2):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, col).value == t)
tot = fila("TOTAL B"); ult = tot - 1
assert ws.cell(ult, 1).value == "B.8.2", ws.cell(ult, 1).value
sub = fila("B.8 Sector hall y circulación")
N = 6
# correr referencias >= tot
for row in ws.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            partes = re.split(r"('[^']*'![$A-Z0-9:]+)", c.value)
            c.value = "".join(p if p.startswith("'") else REF.sub(
                lambda m: f"{m.group(1)}{m.group(2)}{m.group(3)}{int(m.group(4)) + N if int(m.group(4)) >= tot else m.group(4)}", p) for p in partes)
alturas = {r: d.height for r, d in ws.row_dimensions.items()}
ws.insert_rows(tot, N)
for r in sorted(alturas, reverse=True):
    if r >= tot: ws.row_dimensions[r + N].height = alturas[r]
cf = ws.conditional_formatting; nuevas = []
for rng in list(cf):
    s = str(rng.sqref); m = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", s)
    if m and int(m.group(4)) >= tot:
        nuevas.append((f"{m.group(1)}{m.group(2)}:{m.group(3)}{int(m.group(4)) + N}", rng.rules))
        del cf._cf_rules[rng]
for s, rules in nuevas:
    for rule in rules: cf.add(s, rule)

ITEMS = [
 ("B.9.1", "Preparación de superficies de pilares y cielo (limpieza, raspado, reparación de fisuras e imprimante)", "m²", "=2*4*0.3*3.5+2.6*4.6",
  "Igual que la pintura de pilares (B.9.2) y de cielo (B.9.3): 8,40 + 11,96 = 20,36 m²."),
 ("B.9.2", "Pintura de pilares", "m²", "=2*4*0.3*3.5",
  "2 pilares de 0,30 × 0,30 m (plano San Miguel v5) × 4 caras × 3,50 m de alto (altura por piso acordada) = 8,40 m²."),
 ("B.9.3", "Pintura de cielo del pórtico", "m²", "=2.6*4.6",
  "Cara inferior de la losa del pórtico, la misma superficie de alero de F.2.1 (línea «proy. alero techumbre» del plano): 2,60 × 4,60 = 11,96 m²."),
 ("B.9.4", "Preparación y pintura de muretes y remates", "ml", "=1.85+0.65",
  "Murete del pórtico a escala en el plano: del pilar al muro de la casona 1,85 ml + retorno junto al pilar 0,65 ml = 2,50 ml. Por metro lineal: ambas caras y el remate."),
 ("B.9.5", "Preparación y pintura de vigas", "ml", "=3.4+1.85+1.85",
  "Vigas del pórtico a escala en el plano: entre los 2 pilares 3,40 ml + de cada pilar al muro de la casona 1,85 × 2 = 7,10 ml. Por metro lineal: la sección de la viga no figura en el plano."),
]
est_sub = [copy.copy(ws.cell(sub, k)._style) for k in range(1, 8)]
est_it = [copy.copy(ws.cell(ult, k)._style) for k in range(1, 8)]
r = tot
for k in range(7): ws.cell(r, k + 1)._style = copy.copy(est_sub[k])
ws.cell(r, 2).value = "B.9 Sector pórtico de acceso"
ws.row_dimensions[r].height = alturas.get(sub)
for cod, des, uni, cant, obs in ITEMS:
    r += 1
    for k in range(7): ws.cell(r, k + 1)._style = copy.copy(est_it[k])
    ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value = cod, des, uni, cant
    ws.cell(r, 6).value = f"=D{r}*E{r}"; ws.cell(r, 7).value = obs
    ws.row_dimensions[r].height = alturas.get(ult)
t = r + 1
assert ws.cell(t, 2).value == "TOTAL B", ws.cell(t, 2).value
ws.cell(t, 6).value = re.sub(r"F(\d+)\)$", f"F{t-1})", ws.cell(t, 6).value)
wb.save(dst); print("ok TOTAL B", ws.cell(t, 6).value, "CF", [str(x.sqref) for x in ws.conditional_formatting])
