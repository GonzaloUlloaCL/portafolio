import sys, importlib.util
spec = importlib.util.spec_from_file_location("v", sys.argv[2])
# reutiliza el evaluador del verificador sin ejecutar sus chequeos
src = open(sys.argv[2], encoding="utf-8").read().split("# ------------------------------------------------------------ lectura de la estructura")[0]
ns = {}
exec(compile(src.replace("ARCHIVO, PLANO, BORRADOR = sys.argv[1], sys.argv[2], sys.argv[3]", ""), "v", "exec"), ns)
import openpyxl
wb = openpyxl.load_workbook(sys.argv[1]); ws = wb["Sede San Miguel"]; cache = {}
mostrar = set(sys.argv[3].split(","))
letra = None
for f in range(1, ws.max_row + 1):
    a, b = ws[f"A{f}"].value, ws[f"B{f}"].value
    if isinstance(a, str) and len(a) == 1 and ws[f"C{f}"].value == "Unidad":
        letra = a
        if letra in mostrar: print(f"\n[{a}] {b}")
        continue
    if letra in mostrar and (a is not None or b):
        d = ws[f"D{f}"].value
        dv = ns["evaluar"](wb, "Sede San Miguel", f"D{f}", cache) if d is not None else None
        txt = f"{dv:8.2f}" if dv is not None else "       ·"
        print(f"  {str(a or ''):5} {str(b or '')[:62]:62} {str(ws[f'C{f}'].value or ''):3} {txt}  {str(ws[f'H{f}'].value or ''):10} {str(ws[f'I{f}'].value or '')[:70]}")
