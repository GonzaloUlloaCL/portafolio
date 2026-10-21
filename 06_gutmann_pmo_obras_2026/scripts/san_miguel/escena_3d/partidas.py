"""Exporta las partidas del presupuesto sala por sala a JSON (cantidades evaluadas, sin precios)."""
import sys, json
src = open(sys.argv[2], encoding="utf-8").read().split("# ------------------------------------------------------------ lectura de la estructura")[0]
ns = {}
exec(compile(src.replace("ARCHIVO, PLANO, BORRADOR = sys.argv[1], sys.argv[2], sys.argv[3]", ""), "v", "exec"), ns)
import openpyxl
wb = openpyxl.load_workbook(sys.argv[1]); ws = wb["Sede San Miguel"]; cache = {}
sec = {}; letra = None
for f in range(1, ws.max_row + 1):
    a, b = ws[f"A{f}"].value, ws[f"B{f}"].value
    if isinstance(a, str) and len(a) == 1 and ws[f"C{f}"].value == "Unidad":
        letra = a; sec[a] = {"titulo": b, "items": []}; continue
    if letra and a is not None and b and "TOTAL" not in str(b):
        d = ws[f"D{f}"].value
        q = ns["evaluar"](wb, "Sede San Miguel", f"D{f}", cache) if d is not None else None
        sec[letra]["items"].append([str(a), b, ws[f"C{f}"].value or "", None if q is None else round(q, 2), ws[f"H{f}"].value or ""])
    if isinstance(b, str) and b.startswith("RESUMEN"): break
json.dump(sec, open(sys.argv[3], "w"), ensure_ascii=False)
for k, v in sec.items(): print(k, v["titulo"], len([i for i in v["items"] if "." in i[0]]))
