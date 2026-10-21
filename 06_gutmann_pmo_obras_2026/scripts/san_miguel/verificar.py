"""Verifica el presupuesto sala por sala sobre el ARCHIVO FINAL (no sobre el generador).

Uso: python -I verificar.py <presupuesto.xlsx> <plano.pdf> <borrador_7oct.xlsx>
Sale con código 1 si algún chequeo falla. Incluye controles negativos: rompe una copia en memoria
y exige que los chequeos detecten la rotura.
"""
import copy
import re
import sys
import unicodedata

import openpyxl
import pymupdf

ARCHIVO, PLANO, BORRADOR = sys.argv[1], sys.argv[2], sys.argv[3]
HOJA = "Sede San Miguel"
fallas = []


def ok(cond, msg):
    print(("  OK   " if cond else "  FALLA ") + msg)
    if not cond:
        fallas.append(msg)
    return cond


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


# ------------------------------------------------------------ evaluador de fórmulas (solo lo que usa el archivo)
REF = re.compile(r"(?:'([^']+)'!)?\$?([A-Z]{1,3})\$?(\d+)")
SUMA = re.compile(r"SUM\(((?:'[^']+'!)?\$?[A-Z]{1,3}\$?\d+):\$?([A-Z]{1,3})\$?(\d+)\)")


def col_num(col):
    k = 0
    for ch in col:
        k = k * 26 + ord(ch) - 64
    return k


def col_letra(k):
    s = ""
    while k:
        k, r = divmod(k - 1, 26)
        s = chr(65 + r) + s
    return s


def evaluar(wb, hoja, coord, cache):
    clave = (hoja, coord)
    if clave in cache:
        return cache[clave]
    v = wb[hoja][coord].value
    if v is None or v == "":
        r = 0.0
    elif isinstance(v, (int, float)):
        r = float(v)
    elif isinstance(v, str) and v.startswith("="):
        r = evaluar_formula(wb, hoja, v[1:], cache)
    else:
        r = 0.0
    cache[clave] = r
    return r


def evaluar_formula(wb, hoja, f, cache):
    def suma(m):
        ini = m.group(1)
        mi = REF.fullmatch(ini)
        h = mi.group(1) or hoja
        c1, r1, c2, r2 = col_num(mi.group(2)), int(mi.group(3)), col_num(m.group(2)), int(m.group(3))
        total = sum(evaluar(wb, h, f"{col_letra(c)}{rr}", cache)
                    for c in range(c1, c2 + 1) for rr in range(r1, r2 + 1))
        return f"({total!r})"

    f = SUMA.sub(suma, f)
    f = REF.sub(lambda m: f"({evaluar(wb, m.group(1) or hoja, m.group(2) + m.group(3), cache)!r})", f)
    if not re.fullmatch(r"[0-9eE+\-*/(). ]*", f):
        raise ValueError(f"fórmula no soportada: {f}")
    return float(eval(f, {"__builtins__": {}}, {}))


# ------------------------------------------------------------ lectura de la estructura
def estructura(ws):
    secciones, actual = [], None
    for fila in range(1, ws.max_row + 1):
        a, b, c = ws[f"A{fila}"].value, ws[f"B{fila}"].value, ws[f"C{fila}"].value
        if isinstance(a, str) and re.fullmatch(r"[A-Z]", a) and c == "Unidad":
            actual = {"letra": a, "titulo": b, "cab": fila, "items": [], "total": None}
            secciones.append(actual)
        elif actual and isinstance(a, str) and re.fullmatch(r"\d+\.\d+", a):
            actual["items"].append(fila)
        elif actual and isinstance(b, str) and b.startswith("TOTAL ") and actual["total"] is None:
            actual["total"] = fila
            actual = None
    fila_cd = next(f for f in range(1, ws.max_row + 1) if ws[f"B{f}"].value == "COSTO DIRECTO")
    filas_res = {ws[f"B{f}"].value: f for f in range(fila_cd, ws.max_row + 1) if ws[f"B{f}"].value}
    return secciones, fila_cd, filas_res


def chequeo_sumas(ws, secciones, fila_cd, silencioso=False):
    """Devuelve la lista de problemas de estructura de sumas (vacía = todo bien)."""
    problemas = []
    todas = {f for s in secciones for f in s["items"]}
    for s in secciones:
        m = re.fullmatch(r"=SUM\(F(\d+):F(\d+)\)", str(ws[f"F{s['total']}"].value))
        if not m:
            problemas.append(f"TOTAL {s['letra']} no es una SUM simple")
            continue
        a, b = int(m.group(1)), int(m.group(2))
        cubiertas = {f for f in todas if a <= f <= b}
        if cubiertas != set(s["items"]):
            problemas.append(f"TOTAL {s['letra']} no cubre exactamente sus partidas")
        for f in s["items"]:
            if ws[f"F{f}"].value != f"=D{f}*E{f}":
                problemas.append(f"partida {s['letra']} fila {f}: total no es =D*E")
    refs = re.findall(r"F(\d+)", str(ws[f"F{fila_cd}"].value))
    esperadas = [str(s["total"]) for s in secciones]
    if sorted(refs) != sorted(esperadas) or len(refs) != len(set(refs)):
        problemas.append("COSTO DIRECTO no suma cada TOTAL exactamente una vez")
    return problemas


wb = openpyxl.load_workbook(ARCHIVO)
ws = wb[HOJA]
secciones, fila_cd, filas_res = estructura(ws)

print("1-3. Estructura de sumas")
ok(len(secciones) == 20, f"20 secciones (hay {len(secciones)})")
ok(all(s["total"] for s in secciones), "cada sección tiene su TOTAL")
problemas = chequeo_sumas(ws, secciones, fila_cd)
ok(not problemas, "TOTAL = SUM exacta de sus partidas; cada partida = cantidad × precio; costo directo suma los 20 "
   "totales una vez" + ("" if not problemas else f" → {problemas[:3]}"))

print("2b. Controles negativos (la verificación tiene que detectar la rotura)")
wb_roto = copy.deepcopy(wb)
ws_roto = wb_roto[HOJA]
ws_roto[f"F{fila_cd}"].value = ws_roto[f"F{fila_cd}"].value.replace(f"+F{secciones[5]['total']}", "")
ok(any("COSTO DIRECTO" in p for p in chequeo_sumas(ws_roto, secciones, fila_cd)),
   f"quitar el TOTAL {secciones[5]['letra']} del costo directo se detecta")
wb_roto = copy.deepcopy(wb)
ws_roto = wb_roto[HOJA]
f0 = secciones[3]["items"][2]
ws_roto[f"F{f0}"].value = f"=D{f0}"
ok(any("=D*E" in p for p in chequeo_sumas(ws_roto, secciones, fila_cd)), "una partida con total mal escrito se detecta")
wb_roto = copy.deepcopy(wb)
ws_roto = wb_roto[HOJA]
s0 = secciones[7]
ws_roto[f"F{s0['total']}"].value = f"=SUM(F{s0['items'][0]}:F{s0['items'][-2]})"
ok(any(f"TOTAL {s0['letra']}" in p for p in chequeo_sumas(ws_roto, secciones, fila_cd)),
   "un TOTAL que deja fuera su última partida se detecta")

print("3b. Cálculo con precios de prueba ($1.000 en todas las partidas, en una copia)")
wb_p = copy.deepcopy(wb)
ws_p = wb_p[HOJA]
items = [f for s in secciones for f in s["items"]]
for f in items:
    ws_p[f"E{f}"].value = 1000
cache = {}
cantidades = {f: evaluar(wb_p, HOJA, f"D{f}", cache) for f in items}
esperado_cd = sum(cantidades.values()) * 1000
gg, ut, iva = (ws[f"E{filas_res[k]}"].value for k in ("GASTOS GENERALES", "UTILIDADES", "IVA"))
esperado_total = esperado_cd * (1 + gg + ut) * (1 + iva)
total_hoja = evaluar(wb_p, HOJA, f"F{filas_res['TOTAL']}", cache)
cd_hoja = evaluar(wb_p, HOJA, f"F{fila_cd}", cache)
ok(abs(cd_hoja - esperado_cd) < 0.01, f"costo directo de la hoja {cd_hoja:,.0f} = esperado {esperado_cd:,.0f}")
ok(abs(total_hoja - esperado_total) < 0.01, f"TOTAL de la hoja {total_hoja:,.0f} = esperado {esperado_total:,.0f}")
ok((gg, ut, iva) == (0.05, 0.15, 0.19), f"porcentajes GG {gg}, utilidades {ut}, IVA {iva}")
ws_p[f"F{fila_cd}"].value = ws_p[f"F{fila_cd}"].value.replace(f"+F{secciones[9]['total']}", "")
total_roto = evaluar(wb_p, HOJA, f"F{filas_res['TOTAL']}", {})
ok(abs(total_roto - esperado_total) > 1, "control negativo: sin un TOTAL en el costo directo, el total ya no calza")

print("4. Recintos y superficies contra el PDF")
texto = pymupdf.open(PLANO)[0].get_text()
rotulos = set(re.findall(r"Sala de Capacitación \d{2}|Oficina \d{2}|Recepción|Estar de Alumnos|"
                         r"B\. (?:Mujeres|Hombres) \d|B\. Universal", texto))
supers = dict(re.findall(r"(Sala de Capacitación \d{2}|Oficina \d{2}|Recepción|Estar de Alumnos)\s*Sup\. útil "
                         r"(\d+,\d+)", texto))


def recinto_de(titulo):
    t = re.sub(r"\s*\(.*\)$", "", titulo)
    return norm(t).replace("bano ", "b. ")


titulos = {recinto_de(s["titulo"]) for s in secciones}
faltan = {r for r in rotulos if norm(r) not in titulos}
ok(len(rotulos) == 15 and not faltan, f"los 15 rótulos del plano tienen su sección (faltan: {sorted(faltan) or 'ninguno'})")
cb = wb["Cubicación base"]
filas_cb = {cb[f"A{f}"].value: f for f in range(5, 40) if cb[f"A{f}"].value}
dif = []
for nombre, sup in supers.items():
    v = cb[f"C{filas_cb[nombre]}"].value
    if abs(v - float(sup.replace(",", "."))) > 1e-9:
        dif.append((nombre, v, sup))
ok(len(supers) == 10 and not dif, f"las 10 superficies rotuladas coinciden en «Cubicación base» {dif or ''}")
cache = {}
for nombre in supers:
    f = filas_cb[nombre]
    if cb[f"G{f}"].value is not None:
        calc = evaluar(wb, "Cubicación base", f"G{f}", cache)
        rot = cb[f"C{f}"].value
        ok(abs(calc / rot - 1) <= 0.03, f"{nombre}: calculada {calc:.2f} vs rotulada {rot:.2f} ({calc / rot - 1:+.1%})")

print("4b. Alturas de piso a cielo de la casona (las midió Gonza el 08-oct) y muros a pintar")
ALTURAS = {"recepcion (sin el bow window)": 3.36, "bow window": 3.47, "hall": 3.30, "pasillo de los banos 1": 2.38,
           "estar de alumnos": 3.40, "oficina 01": 3.30, "banos 1: cubiculos": 3.37,
           "b. hombres 2: zona de sanitarios": 3.40, "b. hombres 2: zona de lavamanos": 3.30,
           "sala pegada al estar, al otro lado del vano nuevo": 2.50}
filas_mur, en_tabla = {}, False
for f in range(1, cb.max_row + 1):
    a = norm(cb[f"A{f}"].value)
    if a == "zona":
        en_tabla = True
    elif en_tabla and a.startswith("total muros"):
        break
    elif en_tabla:
        filas_mur[a] = f
leidas = {k: cb[f"C{f}"].value for k, f in filas_mur.items()}
ok(leidas == ALTURAS, f"las 10 alturas de Gonza están en «Cubicación base» {'' if leidas == ALTURAS else leidas}")
cache = {}
malos = [f for f in filas_mur.values() if cb[f"D{f}"].value is not None
         and abs(evaluar(wb, "Cubicación base", f"F{f}", cache) - cb[f"C{f}"].value * cb[f"D{f}"].value) > 1e-9]
ok(not malos, f"sup. de muros = altura × largo en cada zona {malos or ''}")
pint = [f for s in secciones if s["letra"] in "JKLM" for f in s["items"] if norm(ws[f"B{f}"].value) == "pintura de muros"]
ok(len(pint) == 4 and all("Cubicación base" in str(ws[f"D{f}"].value) for f in pint),
   "la pintura de muros de la Recepción, el hall, el Estar y la Oficina 01 sale de esa tabla")

print("5. Ítems de la plantilla de Luciano y adicional")
sec = {s["letra"]: s for s in secciones}
desc_a = [norm(ws[f"B{f}"].value) for f in sec["A"]["items"]]
for clave in ("cierre perimetral provisorio", "oficina tecnica y de inspeccion", "bodega y panol de herramientas",
              "comedor y vestidores para personal", "servicios higienicos y duchas", "empalmes provisorios",
              "letrero de obra y senaletica"):
    ok(any(clave in d for d in desc_a), f"trabajos previos: «{clave}»")
t = sec["T"]["items"]
ok(len(t) == 1 and ws[f"D{t[0]}"].value == 5 and "perforaciones" in norm(ws[f"B{t[0]}"].value)
   and ws[f"H{t[0]}"].value == "Confirmado", "adicional del frontis: 5 perforaciones, Confirmado")

print("6. Coherencia Estado / cantidad")
malos = []
for f in items:
    est, cant, obs = ws[f"H{f}"].value, ws[f"D{f}"].value, ws[f"I{f}"].value
    if est == "Pendiente" and (cant is not None or not obs):
        malos.append((f, "pendiente con cantidad o sin observación"))
    if est != "Pendiente" and cant is None:
        malos.append((f, f"{est} sin cantidad"))
    if est == "Global" and cant != 1:
        malos.append((f, "global con cantidad distinta de 1"))
ok(not malos, f"Pendiente = sin cantidad y con observación; el resto tiene cantidad {malos[:3] or ''}")

print("7. Trazabilidad con el borrador del 7-oct")
MAPA = {
    "instalacion de faena (oficina, bodega, panol)": [["oficina tecnica"], ["bodega y panol"]],
    "cierre perimetral provisorio": [["cierre perimetral provisorio"]],
    "empalmes provisorios de luz y agua": [["empalmes provisorios"]],
    "servicios higienicos provisorios": [["servicios higienicos"]],
    "letrero de obra y senaletica": [["letrero de obra"]],
    "tala y retiro de arboles": [["tala"]],
    "desarme y retiro de escombros": [["demolicion"], ["retiro de escombros"]],
    "trazado, replanteo y niveles": [["trazado"]],
    "relleno de piso": [["relleno"]],
    "losa radier": [["radier"]],
    "fundaciones y sobrecimientos de obra nueva": [["fundaciones y sobrecimientos"]],
    "muros de albanileria nuevos (h = 3,50 m)": [["muros nuevos de albanileria"]],
    "refuerzo de acero: pilares 100x100x3 mm": [["pilares de acero 100"]],
    "refuerzo de acero: vigas 200x50x3 mm": [["viga de acero 200"]],
    "tabiqueria interior": [["tabiques nuevos"]],
    "pavimento (material por confirmar)": [["pavimento (material"]],
    "zocalo / guardapolvo": [["guardapolvo"]],
    "cielo (material por confirmar)": [["cielo (material"]],
    "pintura de muros": [["pintura de muros"]],
    "pintura de cielo": [["pintura de cielo"]],
    "puertas": [["puerta"]],
    "ventanas y vidrios": [["ventana"]],
    "instalacion electrica e iluminacion": [["instalacion electrica"]],
    "pavimento (porcelanato o ceramica por confirmar)": [["pavimento ceramico"]],
    "revestimiento de muros": [["revestimiento de muros"]],
    "cielo y pintura": [["cielo y pintura"]],
    "puertas y cubiculos": [["puerta de acceso"], ["puertas de cubiculo"]],
    "artefactos sanitarios (wc, lavamanos, urinarios)": [["wc"], ["lavamanos"], ["urinarios"]],
    "griferia": [["griferia"]],
    "accesorios de bano": [["accesorios de bano"]],
    "instalacion sanitaria y ventilacion forzada": [["instalacion sanitaria"], ["shaft de extraccion"]],
    "techumbre salas": [["techumbre"]],
    "aleros": [["aleros"]],
    "aguas lluvia: canales y bajadas": [["aguas lluvia"]],
    "fachada (revestimiento / reparacion)": [["fachada"]],
    "barandas h = 0,95 m": [["barandas"]],
    "perimetro: reja exterior": [["reja exterior"]],
    "porton": [["porton"]],
    "muro colindante": [["muro colindante"]],
    "pavimentos exteriores y circulaciones": [["pavimento de pasillos"], ["areas exteriores"]],
    "tablero electrico general y empalme definitivo": [["tablero electrico general"]],
    "red de agua potable": [["agua potable"]],
    "red de alcantarillado": [["alcantarillado"]],
    "aseo durante la obra": [["aseo durante la obra"]],
    "aseo final": [["aseo final"]],
}
DESCARTES = {
    "casona segundo piso (todas las partidas)": "Gonza, 08-oct: el piso 2 se agrega cuando tenga medidas",
    "climatizacion y ventilacion": "no aparece en el plano: pasa a «Pendientes» como pregunta",
}
viejo = openpyxl.load_workbook(BORRADOR)["Presupuesto"]
descs_viejas = sorted({norm(str(viejo[f"B{f}"].value).replace("×", "x")) for f in range(1, viejo.max_row + 1)
                       if re.fullmatch(r"\d+(\.\d+)+", str(viejo[f"A{f}"].value or ""))})
descs_nuevas = [norm(str(ws[f"B{f}"].value).replace("×", "x")) for f in items]
sin_destino = []
for d in descs_viejas:
    if d in DESCARTES:
        print(f"  --   descarte: «{d}» ({DESCARTES[d]})")
        continue
    reglas = MAPA.get(d)
    if reglas is None or not all(any(all(k in n_ for k in grupo) for n_ in descs_nuevas) for grupo in reglas):
        sin_destino.append(d)
ok(not sin_destino, f"las {len(descs_viejas)} partidas distintas del borrador tienen destino o descarte explicado "
   f"{sin_destino or ''}")
ok(any("climatizacion" in norm(wb["Pendientes"][f"B{f}"].value) for f in range(4, 30)),
   "climatización quedó como pregunta en «Pendientes»")

print("8. Formato")
malos = []
for hoja in wb.worksheets:
    for fila in hoja.iter_rows():
        for c in fila:
            fl = c.fill
            if fl is not None and fl.fill_type == "solid":
                s, e = fl.start_color.rgb, fl.end_color.rgb
                if not (isinstance(s, str) and s == e and s.upper().startswith("FF")):
                    malos.append(f"{hoja.title}!{c.coordinate}")
ok(not malos, f"todos los rellenos son opacos con inicio = fin (Google Sheets) {malos[:5] or ''}")
reglas_cf = [(str(r.sqref), cf.formula) for r in ws.conditional_formatting for cf in r.rules]
ok(any(sq == f"D{items[0]}:D{items[-1]}" for sq, _ in reglas_cf), f"formato condicional en D{items[0]}:D{items[-1]}")
ok(str(ws.print_area).endswith(f"$A$1:$F${filas_res['TOTAL']}"), f"área de impresión {ws.print_area}")

print("\nResumen de cantidades:")
cuenta = {}
for f in items:
    cuenta[ws[f"H{f}"].value] = cuenta.get(ws[f"H{f}"].value, 0) + 1
print("  partidas:", len(items), "·", ", ".join(f"{k} {v}" for k, v in sorted(cuenta.items())))
cache = {}
for clave, col in (("TOTAL superficie interior (todos los recintos)", "P"),
                   ("TOTAL casona: piso que se retira (Recepción a B. Hombres 2)", "P")):
    print(f"  {clave}: {evaluar(wb, 'Cubicación base', f'{col}{filas_cb[clave]}', cache):.2f} m²")
fila_tot_mur = max(filas_mur.values()) + 1
print(f"  {cb[f'A{fila_tot_mur}'].value}: {evaluar(wb, 'Cubicación base', f'F{fila_tot_mur}', cache):.2f} m²")
print("\n" + ("TODO OK" if not fallas else f"{len(fallas)} FALLA(S)"))
sys.exit(1 if fallas else 0)
