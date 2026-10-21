"""Arma el presupuesto de la ampliación San Bernardo con el formato del de San Miguel v2.0.

Uso: python -I armar_presupuesto.py <presupuesto_san_miguel_v2.0.xlsx> <salida.xlsx>
Toma de San Miguel solo los estilos y el encabezado; las partidas salen de partidas.py.
Nunca uses como salida un Excel que alguien ya editó a mano: genera a otro nombre y traspasa.
"""
import os
import sys
from copy import copy

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from partidas import GG, IVA, SECCIONES, UTILIDAD, totales, valor  # noqa: E402

BASE, SALIDA = sys.argv[1], sys.argv[2]
FECHA = "14-10-2026"

wb = openpyxl.load_workbook(BASE)
ws = wb["Presupuesto"]
# estilos de muestra, tomados antes de borrar
MUESTRA = {k: [copy(c._style) for c in ws[f][:8]] for k, f in
           dict(cabecera=7, sector=8, partida=9, subtotal=17, total=18).items()}
alto_partida = ws.row_dimensions[9].height
for n in [s for s in wb.sheetnames if s != "Presupuesto"]:
    del wb[n]
ws.delete_rows(7, ws.max_row)
ws.conditional_formatting = type(ws.conditional_formatting)()
ws["A4"] = "OBRA: PROPUESTA DE AMPLIACIÓN SAN BERNARDO"
ws["A5"] = "DIRECCIÓN: San Bernardo, RM (por confirmar)"
ws["A6"] = f"FECHA: {FECHA}"

fila = 7


def estilo(r, tipo):
    for c, s in zip(ws[r][:8], MUESTRA[tipo]):
        c._style = copy(s)


subtotales_sec = {}
for letra, nombre, sectores in SECCIONES:
    estilo(fila, "cabecera")
    ws.row_dimensions[fila].height = 15.75
    ws.cell(fila, 2, f"{letra} {nombre}")
    for col, t in zip("CDEFG", ["Unidad", "Cantidad", "P. Unitario", "Total", "Observación"]):
        ws[f"{col}{fila}"] = t
    fila += 1
    filas_sub = []
    for sector, partidas in sectores:
        estilo(fila, "sector")
        ws.cell(fila, 2, sector)
        fila += 1
        ini = fila
        for cod, desc, un, q, pu, obs in partidas:
            estilo(fila, "partida")
            ws.row_dimensions[fila].height = alto_partida
            ws.cell(fila, 1, cod)
            ws.cell(fila, 2, desc)
            ws.cell(fila, 3, un)
            ws.cell(fila, 4, f"={q}" if isinstance(q, str) else q)
            ws.cell(fila, 5, pu)
            ws.cell(fila, 6, f"=D{fila}*E{fila}")
            ws.cell(fila, 7, obs)
            fila += 1
        estilo(fila, "subtotal")
        ws.cell(fila, 2, f"Subtotal {sector.split(' ')[0]}")
        ws.cell(fila, 6, f"=SUM(F{ini}:F{fila - 1})")
        filas_sub.append(fila)
        fila += 1
    estilo(fila, "total")
    ws.cell(fila, 2, f"TOTAL {letra}")
    ws.cell(fila, 6, "=" + "+".join(f"F{r}" for r in filas_sub))
    subtotales_sec[letra] = fila
    fila += 2
ultima_partida = fila - 2

ws.conditional_formatting.add(f"D8:D{ultima_partida}", FormulaRule(
    formula=[f"AND(LEN($A8)>4,ISBLANK($D8))"], fill=PatternFill("solid", start_color="FFF4CCCC", end_color="FFF4CCCC")))

# resumen, igual que San Miguel
estilo(fila, "total")
ws.cell(fila, 2, "RESUMEN")
ws.cell(fila, 6, None)
fila += 1
estilo(fila, "total")
ws.cell(fila, 2, "COSTO DIRECTO")
ws.cell(fila, 6, "=" + "+".join(f"F{r}" for r in subtotales_sec.values()))
f_cd = fila
amarillo = MUESTRA["partida"][4]
for etiqueta, pct in (("GASTOS GENERALES", GG), ("UTILIDADES", UTILIDAD)):
    fila += 2
    ws.cell(fila, 2, etiqueta)
    c = ws.cell(fila, 5, pct)
    c._style = copy(amarillo)
    c.number_format = "0%"
    ws.cell(fila, 6, f"=F{f_cd}*E{fila}").number_format = "#,##0"
    if etiqueta.startswith("GASTOS"):
        f_gg = fila
    else:
        f_ut = fila
fila += 2
ws.cell(fila, 2, "SUBTOTAL NETO")
ws.cell(fila, 6, f"=F{f_cd}+F{f_gg}+F{f_ut}").number_format = "#,##0"
f_neto = fila
fila += 2
ws.cell(fila, 2, "IVA")
c = ws.cell(fila, 5, IVA)
c._style = copy(amarillo)
c.number_format = "0%"
ws.cell(fila, 6, f"=F{f_neto}*E{fila}").number_format = "#,##0"
f_iva = fila
fila += 2
ws.cell(fila, 2, "TOTAL").font = Font(name="Calibri", size=10, bold=True)
c = ws.cell(fila, 6, f"=F{f_neto}+F{f_iva}")
c.number_format = "#,##0"
c.font = Font(name="Calibri", size=10, bold=True)

# ------------------------------------------------------------------ hojas de apoyo
AZUL = PatternFill("solid", start_color="FF1F3A5F", end_color="FF1F3A5F")
CELESTE = PatternFill("solid", start_color="FFE7ECF3", end_color="FFE7ECF3")
BLANCA_B = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
NORMAL = Font(name="Calibri", size=10)
NEGRITA = Font(name="Calibri", size=10, bold=True)
TITULO = Font(name="Calibri", size=12, bold=True)
fino = Side(style="thin", color="FFBFBFBF")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
AJUSTE = Alignment(wrap_text=True, vertical="top")


def hoja(nombre, anchos, indice=None):
    h = wb.create_sheet(nombre, indice)
    for col, w in zip("ABCDEFGHIJ", anchos):
        h.column_dimensions[col].width = w
    return h


def tabla(h, fila0, cabecera, filas):
    for j, t in enumerate(cabecera, 1):
        c = h.cell(fila0, j, t)
        c.fill, c.font, c.alignment, c.border = AZUL, BLANCA_B, AJUSTE, BORDE
    for i, f in enumerate(filas, 1):
        for j, v in enumerate(f, 1):
            c = h.cell(fila0 + i, j, v)
            c.font, c.alignment, c.border = NORMAL, AJUSTE, BORDE
            if isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith("=")):
                c.number_format = "#,##0.00"
    return fila0 + len(filas) + 1


# Léeme
le = hoja("Léeme", [130], 0)
lineas = [
    ("PRESUPUESTO DESGLOSADO POR ÍTEMS A SUMA ALZADA — BORRADOR PARA REVISIÓN", TITULO),
    ("Obra: Propuesta de ampliación San Bernardo · 3 salas de capacitación, guardar, SS.HH., corredor cubierto y patio", None),
    (f"Preparado para Luciano por Gonza el {FECHA}, con la misma base, forma y valores que San Miguel v2.0.", None),
    ("", None),
    ("Cómo está ordenado", NEGRITA),
    ("1. Hoja «Presupuesto», en el orden de la obra: A Obras preliminares, B Recintos (un bloque por recinto), "
     "C Corredor cubierto y patio, D Obras generales (base, radier, instalaciones y aseo). Al final, el resumen con GG 5%, "
     "utilidades 15% e IVA.", None),
    ("2. Las cantidades son fórmulas a la vista (por ejemplo =(10+10)*2.4): la cuenta se ve al hacer clic en la celda, y "
     "la columna «Observación» explica de dónde sale cada número.", None),
    ("3. Celdas amarillas = datos editables (cantidad, precio, porcentajes). Una cantidad vacía se pinta en rojo.", None),
    ("4. Hoja «Para Luciano»: lo que hay que confirmar o decidir antes de ofertar, con el impacto en pesos.", None),
    ("5. Hoja «Cubicación base»: superficies, perímetros, muros y ventanas de cada recinto, con el control contra la "
     "Sup. útil rotulada en el plano.", None),
    ("", None),
    ("De dónde salen los números", NEGRITA),
    ("• Plano «propuesta ampliacion - san bernardo.pdf» (planta, escala 1:50): cotas, superficies útiles rotuladas, "
     "puertas, ventanas, artefactos y pilares. Lo que no tiene cota se midió a escala sobre el dibujo vectorial y dice "
     "«Supuesto» en la observación.", None),
    ("• Correo del arquitecto «Propuesta de Ampliación - San Bernardo»: materiales y sistema constructivo (radier, "
     "tabiques, techumbre, corredor, ventanas, pisos e instalaciones) y altura interior de 2,40 m.", None),
    ("• Precios unitarios: los de San Miguel v2.0 (Drive, 13-10-2026). Puertas: doble 750.000 y simple 400.000 "
     "(indicado por Gonza). La araucaria se tala por indicación de Gonza (ver «Para Luciano», punto 1).", None),
    ("", None),
    ("Qué no incluye", NEGRITA),
    ("Electricidad (el correo la deja al especialista), los 3 equipos de aire acondicionado dibujados en el plano y la "
     "pintura anticorrosiva de la estructura de acero. San Miguel tampoco tenía estas partidas.", None),
]
for i, (t, f) in enumerate(lineas, 1):
    c = le.cell(i, 1, t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = f or NORMAL

# Para Luciano
pl = hoja("Para Luciano", [5, 62, 48, 16, 22, 30])
pl.cell(1, 1, "PARA LUCIANO — lo que hay que confirmar antes de ofertar").font = TITULO
pl.cell(2, 1, "Cada punto dice qué se supuso para poder cerrar el número y cuánto cambia si la respuesta es otra.").font = NORMAL
puntos = [
    ("Araucaria: el correo dice que está protegida por ley «desde las raíces hasta la copa». Si es Araucaria araucana "
     "(Monumento Natural, DS 43 de 1990), cortarla requiere autorización de CONAF. ¿Qué especie es y hay permiso?",
     "Se presupuestó la tala (indicación de Gonza). Si el árbol se queda, no se puede excavar ni hacer radier sobre sus "
     "raíces y hay que rediseñar el patio y el corredor norte.", "−766.667 sin tala, más el rediseño", "A.1.6, A.1.8, C, D"),
    ("Vigas del corredor: el correo dice 100×50×3 y el plano rotula 150×50×3. ¿Cuál va?",
     "Se usó la del correo (6,78 kg/m).", "+450.000 con 150×50×3", "C.1.4"),
    ("Tabique perimetral: además del metalcon y el yeso cartón, lleva por fuera OSB o terciado de 15 mm, omega y "
     "plancha metálica. ¿El P.U. de 85.000 de San Miguel incluye ese revestimiento?",
     "Se dio por incluido.", "Si no, falta una partida por 159,74 m²", "B.1.1, B.2.1, B.3.1, B.4.1, B.5.1"),
    ("Piso de las salas: el correo pide «el mismo tipo de piso existente». ¿Qué piso es?",
     "Se usó el porcelanato de San Miguel (95.000/m²).", "Según el piso", "B.1.5, B.2.5, B.3.4, B.4.4, B.5.5"),
    ("Fuera del presupuesto: electricidad, los 3 equipos de aire acondicionado del plano y la pintura del acero. "
     "¿Se agregan o los cotiza otro?", "No están, como en San Miguel.", "Por cotizar", "—"),
    ("Malla C-92: ¿el P.U. del radier (36.000/m²) la incluye?", "Se dio por incluida, como en San Miguel.",
     "Si no, partida aparte", "D.1.3"),
    ("Obras preliminares: son los valores de San Miguel sin escalar, aunque esta obra es casi 3 veces más chica "
     "(286 m² contra 785). ¿Se ajustan?", "Sin escalar; solo se bajó de 3 a 1 baño provisorio.", "Hasta −2,7 M si se "
     "escalan por superficie", "A.1.1 a A.1.5"),
    ("Rebaje de terreno: ¿el P.U. de 50.000/m³ incluye el retiro del material?",
     "Se dio por incluido.", "Si no: 88,57 m³ × P.U. de retiro", "A.1.8"),
    ("Altura del corredor: el plano no la da.", "Pilares de 2,40, igual que el interior de las salas.",
     "Proporcional a la altura", "C.1.3"),
    ("Cotas que no cuadran en el plano: (a) la cadena de recintos de abajo suma 14,84 y la de vanos 15,00; (b) la cadena "
     "del patio suma 10,30 y la cota total dice 10,10.", "Se usaron 14,84 (recintos y muros) y 10,10 (cota total).",
     "+140.000 si el patio mide 10,30", "C.3.2, A.1.8, D.1"),
    ("SS.HH.: altura del porcelanato de muro y de los cubículos.", "Porcelanato hasta el cielo (2,40); cubículos de 2,00.",
     "Proporcional", "B.5.6, B.5.8"),
    ("Puerta entre el corredor norte y el pasillo existente (junto a la Kitchenette): ¿es nueva o ya existe?",
     "Se presupuestó como nueva.", "−400.000 si existe", "C.4.1"),
    ("Aguas lluvia: en San Miguel, F.2.3 dice 1 × 1.500.000 pero su total es 3.500.000. ¿Cuál es el valor?",
     "Se usó 1.500.000.", "+2.000.000 si es 3.500.000", "D.2.1"),
    ("Dirección exacta de la obra, para el encabezado.", "«San Bernardo, RM (por confirmar)».", "—", "Encabezado"),
]
tabla(pl, 4, ["#", "Pregunta", "Qué se supuso", "Impacto", "Partidas", "Respuesta"],
      [(i, p, s, imp, part, "") for i, (p, s, imp, part) in enumerate(puntos, 1)])

# Cubicación base
cb = hoja("Cubicación base", [24, 11, 11, 15, 13, 11, 14, 13, 13, 46])
cb.cell(1, 1, "CUBICACIÓN BASE — medidas del plano y control contra la Sup. útil rotulada").font = TITULO
f = tabla(cb, 3, ["Recinto", "Largo (m)", "Ancho (m)", "Sup. calculada (m²)", "Sup. rotulada (m²)", "Diferencia",
                  "Perímetro (ml)", "Muro perimetral (ml)", "Muro interior (ml)", "Cómo se midió"], [
    ("Sala de Capacitación 01", 10.00, 5.00, "=B4*C4", 50.00, "=D4/E4-1", "=2*(B4+C4)", 20.00, 10.00,
     "Cota vertical 10,00 (derecha) y horizontal 5,00 (abajo)."),
    ("Sala de Capacitación 02", 8.48, 5.00, "=B5*C5", 42.40, "=D5/E5-1", "=2*(B5+C5)", 16.96, 5.00,
     "Cadena de abajo 0,12 · 6,00 · 0,12 · 8,48 · 0,12 y cota vertical 5,00."),
    ("Sala de Capacitación 03", 6.00, 5.00, "=B6*C6", 30.00, "=D6/E6-1", "=2*(B6+C6)", 17.00, 0,
     "Misma cadena (6,00) y cota vertical 5,00 de la izquierda."),
    ("Guardar (en L)", 5.00, 4.56, "=B7*C7-2.6*3.56", 13.54, "=D7/E7-1", "=2*(B7+C7)", 4.68, 0,
     "Rectángulo 5,00 × 4,56 menos el SS.HH. con sus muros (2,60 × 3,56)."),
    ("SS.HH.", 3.44, 2.48, "=B8*C8", None, None, "=2*(B8+C8)", 7.92, 6.16,
     "Cadena de arriba 0,12 · 2,48 · 0,12 y cota vertical 3,44. No trae rótulo."),
])
for r in range(4, 9):
    cb.cell(r, 6).number_format = "0.0%"
f += 1
cb.cell(f, 1, "Superficies generales (base, radier, rebaje y aseo)").font = NEGRITA
f = tabla(cb, f + 1, ["Zona", "Largo (m)", "Ancho (m)", "Superficie (m²)", "", "", "", "", "", "Cómo se midió"], [
    ("Huella bloque sur (Salas 02 y 03)", 14.84, 5.24, f"=B{f+2}*C{f+2}", "", "", "", "", "",
     "Por fuera de muros: 0,12 + 6,00 + 0,12 + 8,48 + 0,12 = 14,84; 0,12 + 5,00 + 0,12 = 5,24."),
    ("Huella bloque oriente (Sala 01, Guardar, SS.HH.)", 14.76, 5.24, f"=B{f+3}*C{f+3}", "", "", "", "", "",
     "Cota total 20,00 menos el bloque sur (5,24)."),
    ("Corredor norte", 6.28, 2.30, f"=B{f+4}*C{f+4}", "", "", "", "", "", "Cotas 6,28 (oriente) y 2,30 (izquierda)."),
    ("Corredor oriente", 10.10, 2.10, f"=B{f+5}*C{f+5}", "", "", "", "", "",
     "Entre la línea de pilares y el muro de la Sala 01 (2,00 + 0,10)."),
    ("Corredor sur", 9.60, 2.00, f"=B{f+6}*C{f+6}", "", "", "", "", "",
     "Desde el borde poniente hasta el muro de la Sala 01; cota 2,00."),
    ("Patio descubierto", 10.10, 7.50, f"=B{f+7}*C{f+7}", "", "", "", "", "", "Cotas 7,50 y 10,10 (ver «Para Luciano», 10)."),
    ("TOTAL", None, None, f"=SUM(D{f+2}:D{f+7})", "", "", "", "", "", "Base de A.1.8, D.1 y D.3."),
])
f += 1
cb.cell(f, 1, "Pilares del corredor (centro de cada pilar en el plano, en metros desde la esquina del dibujo)").font = NEGRITA
f = tabla(cb, f + 1, ["Pilar", "X (m)", "Y (m)", "Borde", "", "", "", "", "", "Nota"], [
    (1, 9.05, 9.48, "norte", "", "", "", "", "", "Cuadrados de 0,10 × 0,10 del dibujo vectorial."),
    (2, 11.14, 9.48, "norte", "", "", "", "", "", ""),
    (3, 13.23, 9.48, "norte / oriente", "", "", "", "", "", "Esquina."),
    (4, 13.23, 12.04, "oriente", "", "", "", "", "", "Separación 2,55 entre ejes (2,45 + 0,10)."),
    (5, 13.23, 14.59, "oriente", "", "", "", "", "", ""),
    (6, 13.23, 17.14, "oriente", "", "", "", "", "", ""),
    (7, 13.23, 19.69, "oriente / sur", "", "", "", "", "", "Esquina."),
    (8, 10.73, 19.69, "sur", "", "", "", "", "", "Separación 2,50 entre ejes (2,40 + 0,10)."),
    (9, 8.23, 19.69, "sur", "", "", "", "", "", ""),
    (10, 5.73, 19.69, "sur", "", "", "", "", "", "Desde aquí, la viga remata en el muro de la Sala 03 (2,10)."),
])
f += 1
cb.cell(f, 1, "Niveles y paquete de piso").font = NEGRITA
tabla(cb, f + 1, ["Capa", "Espesor (m)", "Nivel (m)", "", "", "", "", "", "", "Fuente"], [
    ("Piso terminado (N.P.T.)", None, 0.10, "", "", "", "", "", "", "Igual al existente (+0,10 en el plano); correo."),
    ("Porcelanato y adhesivo", 0.01, None, "", "", "", "", "", "", "Supuesto."),
    ("Radier H25 con malla C-92", 0.10, None, "", "", "", "", "", "", "Correo."),
    ("Gravilla", 0.10, None, "", "", "", "", "", "", "Correo."),
    ("Base estabilizada", 0.15, None, "", "", "", "", "", "", "Correo («al menos 15 cm»)."),
    ("Fondo de excavación", None, -0.26, "", "", "", "", "", "", "0,10 − 0,36."),
    ("Terreno natural (N.T.N.)", None, 0.05, "", "", "", "", "", "", "Plano: N.T.N. +0,05."),
    ("Rebaje", 0.31, None, "", "", "", "", "", "", "0,05 − (−0,26). Base de A.1.8."),
])

wb.move_sheet("Presupuesto", offset=-(len(wb.sheetnames) - 2))
wb.active = wb.sheetnames.index("Presupuesto")
wb.save(SALIDA)

por_sec, cd, neto, iva, total = totales()
print("guardado:", SALIDA, "· hojas:", wb.sheetnames)
for k, v in por_sec.items():
    print(f"  {k}: {v:,.0f}")
print(f"  CD {cd:,.0f} · neto {neto:,.0f} · IVA {iva:,.0f} · TOTAL {total:,.0f}")
